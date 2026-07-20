from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.domain.models import OfficeEngine
from src.office.backend import OfficeBackend
from src.workbook.preview import inspect_workbook


def warning_cells(path: Path, info) -> list[str]:
    book = load_workbook(path, read_only=True, data_only=False, keep_vba=True, keep_links=False)
    try:
        ws = book[info.selected_sheet]
        cells: set[str] = set()
        for row in range(info.header_row + 1, ws.max_row + 1):
            status = str(ws.cell(row, info.result_columns.status).value or "")
            explanation = str(ws.cell(row, info.result_columns.explanation).value or "")
            if any(token in status for token in ("待确认", "地址为空", "定位失败", "解析失败")):
                cells.add(f"{get_column_letter(info.result_columns.status)}{row}")
                cells.add(f"{get_column_letter(info.result_columns.explanation)}{row}")
            if "多目的地" in status or "目的城市冲突" in explanation or "目的地址为空" in explanation:
                cells.add(f"{get_column_letter(info.recommended_mapping.destination_city)}{row}")
                cells.add(f"{get_column_letter(info.recommended_mapping.destination_address)}{row}")
            if "发货城市冲突" in explanation or "发货地址为空" in explanation:
                cells.add(f"{get_column_letter(info.recommended_mapping.origin_city)}{row}")
                cells.add(f"{get_column_letter(info.recommended_mapping.origin_address)}{row}")
        return sorted(cells)
    finally:
        book.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("engine", choices=("excel", "wps"))
    args = parser.parse_args()
    engine = OfficeEngine(args.engine)
    source = ROOT / "samples" / "input" / "副本26年6月干线账单 2.1-对账2.0-物流商(2).xlsx"
    info = inspect_workbook(source)
    output = ROOT / "samples" / "expected" / "prototype" / (
        "Excel_样表_模拟结果.xlsm" if engine == OfficeEngine.EXCEL else "WPS_样表_模拟结果.xlsm"
    )
    cells = warning_cells(output, info)
    backend = OfficeBackend(engine, ROOT / "temp" / "cf_verify" / engine.value, ROOT / "logs" / "prototype" / "cf_verify")
    response = backend.verify(
        path=output, sheet_name=info.selected_sheet, header_row=info.header_row,
        result_columns=info.result_columns, target_rows=[], last_data_row=info.max_row,
        warning_cells=cells,
    )
    display = response["result"]["conditional_display"]
    report = {
        "engine": engine.value, "output": str(output), "warning_cells": cells,
        "conditional_display": display,
        "all_have_rules": all(item.get("format_condition_count", 0) >= 1 for item in display.values()),
        "display_format_supported": all("display_format_error" not in item for item in display.values()),
        "all_displayed_red": all(
            item.get("display_fill_color") == 13551615 and item.get("display_font_color") == 393372
            for item in display.values()
        ),
        "session": response["session"],
    }
    report["passed"] = (
        report["all_have_rules"] and report["display_format_supported"] and report["all_displayed_red"]
        and report["session"].get("own_pid_residual") is False
    )
    report_path = ROOT / "samples" / "expected" / "prototype" / f"{engine.value}_conditional_format_com.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("engine", "warning_cells", "all_have_rules", "display_format_supported", "all_displayed_red", "passed")}, ensure_ascii=False, indent=2))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
