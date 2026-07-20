from __future__ import annotations

from collections import Counter
from datetime import datetime
import json
from pathlib import Path
import sqlite3
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.utils.files import sha256_file
from src.validation.quote import is_target_quote
from src.workbook.preview import inspect_workbook


OUTPUT = ROOT / "samples" / "expected" / "address_confirmation"
EVIDENCE = ROOT / "docs" / "evidence"
EXCEL = OUTPUT / "Excel_样表_可信地址库结果_TASK-005A_待验收.xlsm"
WPS = OUTPUT / "WPS_样表_可信地址库结果_TASK-005A_待验收.xlsm"
PENDING = ROOT / "outputs" / "task005a" / "待确认地址清单_TASK-005A.xlsx"
SOURCE_HASH = "52011587F8F12A87749087ABEA0ADCE2557874D1B268A55F348E86F8DFA7A951"
TASK004R_HASHES = {
    "Excel_样表_普通驾车距离结果_待验收.xlsm": "898F890C8AEA057319F8352A5D0931250BB2A3725A189809EC76436E585E0327",
    "WPS_样表_普通驾车距离结果_待验收.xlsm": "ECF7DD03B55822D72BB02D7D117245B8B774C7673998DFD088B161DB88F26861",
}


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _result_rows(path: Path, info) -> dict[int, tuple]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=True, keep_links=False)
    try:
        sheet = workbook[info.selected_sheet]
        return {
            row: tuple(sheet.cell(row, column).value for column in info.result_columns.as_dict().values())
            for row in range(info.header_row + 1, info.max_row + 1)
            if is_target_quote(sheet.cell(row, info.recommended_mapping.quote_type).value)
        }
    finally:
        workbook.close()


def main() -> None:
    source = next((ROOT / "samples" / "input").glob("*.xlsx"))
    info = inspect_workbook(source)
    excel_validation = _read_json(EVIDENCE / "TASK005A_EXCEL_VALIDATION.json")
    wps_validation = _read_json(EVIDENCE / "TASK005A_WPS_VALIDATION.json")
    analysis = _read_json(EVIDENCE / "TASK005A_ADDRESS_ANALYSIS.json")
    excel_rows = _result_rows(EXCEL, info)
    wps_rows = _result_rows(WPS, info)

    connection = sqlite3.connect(ROOT / "cache" / "driving_real.sqlite")
    connection.row_factory = sqlite3.Row
    try:
        confirmed_columns = {
            row["name"] for row in connection.execute("PRAGMA table_info(confirmed_addresses)")
        }
        required_columns = {
            "original_city", "original_address", "cleaned_address", "query_address",
            "formatted_address", "province", "city", "district", "street", "number",
            "level", "longitude", "latitude", "confirmation_status", "confirmed_by",
            "confirmed_at", "data_source", "cleaner_version", "geocode_contract_version",
            "address_hash", "geocode_version",
        }
        address_rows = connection.execute(
            "SELECT confirmation_status,risk_level FROM confirmed_addresses"
        ).fetchall()
        versioned_routes = connection.execute(
            """SELECT COUNT(*) FROM route_cache
               WHERE COALESCE(origin_address_id,'')<>''
                 AND COALESCE(destination_address_id,'')<>''
                 AND COALESCE(origin_geocode_version,'')<>''
                 AND COALESCE(destination_geocode_version,'')<>''"""
        ).fetchone()[0]
    finally:
        connection.close()

    pending_book = load_workbook(PENDING, read_only=True, data_only=False, keep_links=False)
    try:
        pending_sheet = pending_book["待确认地址"]
        pending_headers = [pending_sheet.cell(4, col).value for col in range(1, 9)]
        pending_count = sum(
            1 for row in pending_sheet.iter_rows(min_row=5, values_only=True)
            if any(value is not None for value in row)
        )
    finally:
        pending_book.close()

    task004r_dir = ROOT / "samples" / "expected" / "driving_real"
    old_hashes = {
        name: sha256_file(task004r_dir / name)
        for name in TASK004R_HASHES
    }
    address_status_counts = Counter(row["confirmation_status"] for row in address_rows)
    address_risk_counts = Counter(row["risk_level"] for row in address_rows)
    checks = {
        "source_hash_unchanged": sha256_file(source) == SOURCE_HASH,
        "task004r_files_unchanged": old_hashes == TASK004R_HASHES,
        "excel_validation_passed": excel_validation.get("passed") is True,
        "wps_validation_passed": wps_validation.get("passed") is True,
        "excel_wps_results_identical": excel_rows == wps_rows,
        "target_rows_dynamic_74": len(excel_rows) == analysis["target_row_count"] == 74,
        "confirmed_address_schema_complete": required_columns <= confirmed_columns,
        "confirmed_address_count_31": len(address_rows) == analysis["unique_qualified_address_count"] == 31,
        "address_status_counts_match_analysis": dict(address_status_counts) == analysis["confirmation_status_counts"],
        "address_risk_counts_match_analysis": dict(address_risk_counts) == analysis["risk_counts"],
        "versioned_route_cache_present": versioned_routes == 24,
        "pending_list_uses_address_id": pending_headers == [
            "地址ID", "原始城市", "原始详细地址", "高德标准地址", "定位层级",
            "修正查询地址", "是否确认", "确认备注",
        ] and pending_count == 16,
        "row248_still_not_calculated": excel_rows[248][0] is None and excel_rows[248][1] == "多目的地待确认",
        "no_old_low_precision_status": all(item[1] != "查询成功—定位精度较低" for item in excel_rows.values()),
        "all_office_runs_used_zero_http": (
            excel_validation["run_summary"]["actual_http_calls"] == 0
            and wps_validation["run_summary"]["actual_http_calls"] == 0
        ),
    }
    report = {
        "audited_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source": {"path": str(source), "sha256": sha256_file(source)},
        "outputs": {
            "excel": {"path": str(EXCEL), "sha256": sha256_file(EXCEL)},
            "wps": {"path": str(WPS), "sha256": sha256_file(WPS)},
            "pending": {"path": str(PENDING), "sha256": sha256_file(PENDING)},
        },
        "task004r_hashes": old_hashes,
        "address_status_counts": dict(address_status_counts),
        "address_risk_counts": dict(address_risk_counts),
        "versioned_route_count": versioned_routes,
        "checks": checks,
        "passed": all(checks.values()),
    }
    target = EVIDENCE / "TASK005A_FINAL_AUDIT.json"
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"passed": report["passed"], "checks": checks}, ensure_ascii=False, indent=2))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
