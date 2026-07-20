"""TASK-003A reproducible Office compatibility spike.

All workbook writes target copied test files. The fixed source workbook is opened
read-only by validation code and is never passed to a writable Office session.
"""

from __future__ import annotations

import argparse
import gc
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
import traceback
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable
from xml.etree import ElementTree as ET

import psutil
import pythoncom
import win32com.client
import win32gui
import win32process
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "samples" / "input" / "副本26年6月干线账单 2.1-对账2.0-物流商(2).xlsx"
WORKING = ROOT / "samples" / "working" / "office_spike"
EXPECTED = ROOT / "samples" / "expected" / "office_spike"
LOGS = ROOT / "logs" / "office_spike"
FIXTURES = ROOT / "tests" / "fixtures"
BASELINE = EXPECTED / "source_baseline.json"
EXCEL_A = EXPECTED / "Excel_样表转换结果.xlsm"
EXCEL_B = EXPECTED / "Excel_xlsm再次保存结果.xlsm"
WPS_C = EXPECTED / "WPS_样表转换结果.xlsm"
WPS_D = EXPECTED / "WPS_xlsm再次保存结果.xlsm"

Q_HEADER = "高德货车公路距离（公里）"
R_HEADER = "距离查询状态"
S_HEADER = "距离查询说明"
SIM_STATUS = "模拟测试"
SIM_NOTE = "【模拟数据，不可用于正式业务】Office保存兼容性测试。"
MULTI_STATUS = "多目的地待确认"
MULTI_NOTE = "【模拟数据，不可用于正式业务】检测到多个目的地；Office保存兼容性测试，不写入距离。"
TARGET_QUOTE = "非线路报价"
SPECIAL_ROW = 248
XL_MACRO_ENABLED = 52
MsoAutomationSecurityForceDisable = 3


def now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=json_value), encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def file_meta(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": str(path.resolve()),
        "name": path.name,
        "suffix": path.suffix.lower(),
        "size_bytes": stat.st_size,
        "modified_time": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="microseconds"),
        "sha256": sha256_file(path),
    }


def compact_ranges(values: list[int]) -> list[str]:
    values = sorted(set(values))
    if not values:
        return []
    result: list[str] = []
    start = previous = values[0]
    for value in values[1:]:
        if value == previous + 1:
            previous = value
            continue
        result.append(str(start) if start == previous else f"{start}-{previous}")
        start = previous = value
    result.append(str(start) if start == previous else f"{start}-{previous}")
    return result


def stable_digest(items: Any) -> str:
    raw = json.dumps(items, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=json_value)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest().upper()


def serialize_filter_column(item: Any) -> dict[str, Any]:
    result: dict[str, Any] = {"col_id": item.colId, "hidden_button": item.hiddenButton, "show_button": item.showButton}
    filters = getattr(item, "filters", None)
    if filters is not None:
        result["filters"] = list(filters.filter or [])
        result["blank"] = filters.blank
        result["calendar_type"] = filters.calendarType
        result["date_group_items"] = [
            {name: getattr(entry, name, None) for name in ("year", "month", "day", "hour", "minute", "second", "dateTimeGrouping")}
            for entry in (filters.dateGroupItem or [])
        ]
    custom = getattr(item, "customFilters", None)
    if custom is not None:
        result["custom_filters"] = [
            {"operator": entry.operator, "value": entry.val} for entry in (custom.customFilter or [])
        ]
        result["and"] = getattr(custom, "and_", getattr(custom, "_and", getattr(custom, "and", None)))
    dynamic = getattr(item, "dynamicFilter", None)
    if dynamic is not None:
        result["dynamic_filter"] = {"type": dynamic.type, "value": dynamic.val, "max_value": dynamic.maxVal}
    top10 = getattr(item, "top10", None)
    if top10 is not None:
        result["top10"] = {"top": top10.top, "percent": top10.percent, "value": top10.val, "filter_value": top10.filterVal}
    return result


def style_signature(cell: Any) -> dict[str, Any]:
    fill = cell.fill
    font = cell.font
    border = cell.border
    alignment = cell.alignment
    protection = cell.protection
    return {
        "number_format": cell.number_format,
        "font": [font.name, font.sz, font.bold, font.italic, font.underline, font.strike, font.color.type if font.color else None, font.color.rgb if font.color and font.color.type == "rgb" else None],
        "fill": [fill.fill_type, fill.fgColor.type, fill.fgColor.rgb, fill.fgColor.indexed, fill.fgColor.theme, fill.bgColor.type, fill.bgColor.rgb],
        "border": [border.left.style, border.right.style, border.top.style, border.bottom.style, border.diagonal.style],
        "alignment": [alignment.horizontal, alignment.vertical, alignment.text_rotation, alignment.wrap_text, alignment.shrink_to_fit, alignment.indent],
        "protection": [protection.locked, protection.hidden],
    }


def conditional_formatting_snapshot(ws: Any) -> dict[str, Any]:
    entries: list[dict[str, Any]] = []
    count = 0
    for item in ws.conditional_formatting:
        rules = []
        for rule in ws.conditional_formatting[item]:
            count += 1
            rules.append({
                "type": rule.type,
                "operator": getattr(rule, "operator", None),
                "formula": list(getattr(rule, "formula", None) or []),
                "priority": getattr(rule, "priority", None),
                "stop_if_true": getattr(rule, "stopIfTrue", None),
                "dxf_id": getattr(rule, "dxfId", None),
                "dxf_semantic": str(getattr(rule, "dxf", None)),
            })
        entries.append({"range": str(item.sqref), "rules": rules})
    return {"rule_count": count, "range_count": len(entries), "entries": entries}


def sheet_snapshot(ws: Any) -> dict[str, Any]:
    nonempty: list[dict[str, Any]] = []
    formulas: list[dict[str, str]] = []
    value_rows: list[int] = []
    value_cols: list[int] = []
    styles_ap: list[dict[str, Any]] = []
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value is not None:
                is_formula = cell.data_type == "f" or (isinstance(value, str) and value.startswith("="))
                record = {"cell": cell.coordinate, "kind": "formula" if is_formula else "value", "value": json_value(value)}
                nonempty.append(record)
                value_rows.append(cell.row)
                value_cols.append(cell.column)
                if is_formula:
                    formulas.append({"cell": cell.coordinate, "formula": str(value)})
            if cell.column <= 16 and cell.row <= 281:
                styles_ap.append({"cell": cell.coordinate, "style": style_signature(cell)})
    hidden_rows = [index for index, dimension in ws.row_dimensions.items() if dimension.hidden]
    hidden_cols = [str(key) for key, dimension in ws.column_dimensions.items() if dimension.hidden]
    validations = list(ws.data_validations.dataValidation) if ws.data_validations else []
    if value_rows:
        bounds = f"{get_column_letter(min(value_cols))}{min(value_rows)}:{get_column_letter(max(value_cols))}{max(value_rows)}"
    else:
        bounds = None
    functions = Counter()
    for item in formulas:
        formula = item["formula"]
        if re.search(r"(?:^|[^A-Z0-9_.])(?:_xlfn\.)?IF\s*\(", formula, flags=re.I):
            functions["IF"] += 1
        if re.search(r"(?:_xlfn\.)?DISPIMG\s*\(", formula, flags=re.I):
            functions["DISPIMG"] += 1
    qrs = {str(row): [json_value(ws.cell(row, col).value) for col in (17, 18, 19)] for row in range(1, 282)}
    return {
        "title": ws.title,
        "state": ws.sheet_state,
        "actual_value_range": bounds,
        "openpyxl_dimension": ws.calculate_dimension(),
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "nonempty_cell_count": len(nonempty),
        "nonempty_cells": nonempty,
        "nonempty_cells_digest": stable_digest(nonempty),
        "original_ap_value_formula_digest": stable_digest([item for item in nonempty if ord(item["cell"][0].upper()) - 64 <= 16]),
        "original_ap_style_digest": stable_digest(styles_ap),
        "formulas": {
            "count": len(formulas),
            "positions": [item["cell"] for item in formulas],
            "items": formulas,
            "digest": stable_digest(formulas),
            "if_count": functions["IF"],
            "dispimg_count": functions["DISPIMG"],
            "dispimg_positions": [item["cell"] for item in formulas if "DISPIMG" in item["formula"].upper()],
        },
        "conditional_formatting": conditional_formatting_snapshot(ws),
        "auto_filter": {
            "range": ws.auto_filter.ref,
            "columns": [serialize_filter_column(item) for item in (ws.auto_filter.filterColumn or [])],
        },
        "hidden_rows": {"count": len(hidden_rows), "numbers": hidden_rows, "ranges": compact_ranges(hidden_rows)},
        "hidden_columns": {"count": len(hidden_cols), "columns": hidden_cols},
        "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
        "merged_cells": [str(item) for item in ws.merged_cells.ranges],
        "data_validations": {"count": len(validations), "ranges": [str(item.sqref) for item in validations]},
        "tables": [{"name": table.name, "display_name": table.displayName, "range": table.ref} for table in ws.tables.values()],
        "openpyxl_images": len(ws._images),
        "openpyxl_charts": len(ws._charts),
        "qrs_values_rows_1_281": qrs,
    }


def ooxml_snapshot(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        names = sorted(archive.namelist())
        lower = {name.lower(): name for name in names}
        content_types = archive.read("[Content_Types].xml").decode("utf-8", errors="replace")
        media = []
        for name in names:
            if name.lower().startswith("xl/media/") and not name.endswith("/"):
                content = archive.read(name)
                media.append({"name": name, "size_bytes": len(content), "sha256": hashlib.sha256(content).hexdigest().upper()})
        external_relationships = []
        for name in names:
            if not name.lower().endswith(".rels"):
                continue
            try:
                root = ET.fromstring(archive.read(name))
            except ET.ParseError:
                continue
            for rel in root:
                if rel.attrib.get("TargetMode") == "External":
                    external_relationships.append({"part": name, "type": rel.attrib.get("Type"), "target": rel.attrib.get("Target")})
        key_prefixes = ("xl/workbook", "xl/worksheets/", "xl/styles", "xl/sharedstrings", "xl/theme/", "xl/media/", "xl/cellimages", "xl/drawings/", "xl/vba", "xl/activex/", "xl/charts/", "xl/externallinks/")
        key_parts = []
        for name in names:
            if name.lower().startswith(key_prefixes) or name in ("[Content_Types].xml", "_rels/.rels"):
                content = archive.read(name)
                key_parts.append({"name": name, "size_bytes": len(content), "sha256": hashlib.sha256(content).hexdigest().upper()})
        cell_images = [name for name in names if "cellimage" in name.lower()]
        return {
            "zip_test": archive.testzip(),
            "part_count": len(names),
            "parts": names,
            "key_parts": key_parts,
            "content_type_macro_enabled": "macroEnabled" in content_types,
            "workbook_content_type": next((line for line in content_types.split("><") if "workbook" in line.lower()), None),
            "vba_project": "xl/vbaproject.bin" in lower,
            "vba_signature_parts": [name for name in names if "vbasignature" in name.lower()],
            "media": media,
            "cell_images_parts": cell_images,
            "cell_images_xml_exists": any(name.lower() == "xl/cellimages.xml" for name in names),
            "cell_images_relationship_exists": any(name.lower() == "xl/_rels/cellimages.xml.rels" for name in names),
            "active_x_parts": [name for name in names if name.lower().startswith("xl/activex/")],
            "chart_parts": [name for name in names if name.lower().startswith("xl/charts/")],
            "drawing_parts": [name for name in names if name.lower().startswith("xl/drawings/")],
            "external_link_parts": [name for name in names if name.lower().startswith("xl/externallinks/")],
            "external_relationships": external_relationships,
        }


def workbook_snapshot(path: Path) -> dict[str, Any]:
    before = file_meta(path)
    wb = load_workbook(path, read_only=False, data_only=False, keep_links=True, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        sheets = [sheet_snapshot(ws) for ws in wb.worksheets]
        names = []
        for item in wb.defined_names.values():
            names.append({
                "name": item.name,
                "value": item.attr_text,
                "local_sheet_id": item.localSheetId,
                "hidden": item.hidden,
                "function": item.function,
            })
        snapshot = {
            "captured_at": now(),
            "file": before,
            "python_read_only_inspection": {
                "library": "openpyxl",
                "library_version": __import__("openpyxl").__version__,
                "sheet_count": len(wb.worksheets),
                "sheet_order": [ws.title for ws in wb.worksheets],
                "active_sheet": wb.active.title,
                "sheets": sheets,
                "defined_names": names,
                "external_links_count": len(wb._external_links),
                "has_vba_archive": wb.vba_archive is not None,
            },
            "ooxml_zip_inspection": ooxml_snapshot(path),
        }
    finally:
        wb.close()
    after = file_meta(path)
    snapshot["read_only_integrity"] = {"before": before, "after": after, "unchanged": before == after}
    return snapshot


def process_snapshot(names: set[str]) -> list[dict[str, Any]]:
    windows_by_pid: dict[int, list[dict[str, Any]]] = {}

    def enum_window(hwnd: int, _: Any) -> None:
        try:
            _, pid = win32process.GetWindowThreadProcessId(hwnd)
            title = win32gui.GetWindowText(hwnd)
            if win32gui.IsWindowVisible(hwnd) or title:
                windows_by_pid.setdefault(pid, []).append({"hwnd": hwnd, "title": title, "visible": bool(win32gui.IsWindowVisible(hwnd))})
        except Exception:
            pass

    win32gui.EnumWindows(enum_window, None)
    result = []
    for proc in psutil.process_iter(["pid", "name", "create_time", "exe", "cmdline"]):
        try:
            name = (proc.info.get("name") or "").lower()
            if name not in names:
                continue
            result.append({
                "pid": proc.pid,
                "name": name,
                "created_at": datetime.fromtimestamp(proc.info["create_time"]).astimezone().isoformat(timespec="seconds") if proc.info.get("create_time") else None,
                "exe": proc.info.get("exe"),
                "cmdline": proc.info.get("cmdline"),
                "windows": windows_by_pid.get(proc.pid, []),
            })
        except (psutil.NoSuchProcess, psutil.AccessDenied):
            continue
    return sorted(result, key=lambda item: (item["name"], item["pid"]))


def pid_from_hwnd(hwnd: Any) -> int | None:
    try:
        if not hwnd:
            return None
        return int(win32process.GetWindowThreadProcessId(int(hwnd))[1])
    except Exception:
        return None


def process_exists(pid: int | None) -> bool:
    return bool(pid and psutil.pid_exists(pid))


def configure_office(app: Any) -> dict[str, Any]:
    changes = {}
    for name, value in (
        ("AutomationSecurity", MsoAutomationSecurityForceDisable),
        ("EnableEvents", False),
        ("DisplayAlerts", False),
        ("AskToUpdateLinks", False),
        ("ScreenUpdating", False),
    ):
        try:
            setattr(app, name, value)
            changes[name] = {"requested": value, "accepted": True, "readback": json_value(getattr(app, name, None))}
        except Exception as exc:
            changes[name] = {"requested": value, "accepted": False, "error": repr(exc)}
    return changes


def com_cell_value(cell: Any) -> Any:
    value = cell.Value2
    return json_value(value)


def count_formula_matrix(value: Any) -> int:
    if isinstance(value, tuple):
        return sum(count_formula_matrix(item) for item in value)
    return int(isinstance(value, str) and value.startswith("="))


def com_observe_workbook(app: Any, path: Path, engine: str) -> dict[str, Any]:
    wb = None
    observations: dict[str, Any] = {"engine": engine, "path": str(path), "opened": False, "read_only": True}
    try:
        wb = app.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True, IgnoreReadOnlyRecommended=True, AddToMru=False)
        observations["opened"] = True
        observations["file_format"] = json_value(wb.FileFormat)
        observations["has_vba_project"] = json_value(getattr(wb, "HasVBProject", None))
        observations["sheet_count"] = int(wb.Worksheets.Count)
        observations["sheets"] = []
        for index in range(1, int(wb.Worksheets.Count) + 1):
            ws = wb.Worksheets(index)
            dispimg = {}
            for address in ("O38", "O39", "O58", "O78", "O118"):
                cell = ws.Range(address)
                dispimg[address] = {"formula": json_value(cell.Formula), "formula2": json_value(getattr(cell, "Formula2", None)), "text": json_value(cell.Text), "value": com_cell_value(cell)}
            observations["sheets"].append({
                "index": index,
                "name": str(ws.Name),
                "visible": json_value(ws.Visible),
                "used_range": str(ws.UsedRange.Address),
                "auto_filter_mode": bool(ws.AutoFilterMode),
                "filter_mode": bool(ws.FilterMode),
                "hidden_row_count_1_281": sum(bool(ws.Rows(row).Hidden) for row in range(1, 282)),
                "freeze_panes": bool(app.ActiveWindow.FreezePanes) if app.ActiveWindow is not None else None,
                "split_row": json_value(app.ActiveWindow.SplitRow) if app.ActiveWindow is not None else None,
                "split_column": json_value(app.ActiveWindow.SplitColumn) if app.ActiveWindow is not None else None,
                "formula_count": count_formula_matrix(ws.Range("A1:S281").Formula),
                "shape_count": int(ws.Shapes.Count),
                "button_count": int(ws.Buttons().Count),
                "dispimg_cells": dispimg,
            })
            del ws
        observations["defined_names"] = []
        for index in range(1, int(wb.Names.Count) + 1):
            item = wb.Names(index)
            observations["defined_names"].append({"name": str(item.Name), "refers_to": str(item.RefersTo), "visible": bool(item.Visible)})
            del item
    except Exception as exc:
        observations["error"] = repr(exc)
        observations["traceback"] = traceback.format_exc()
    finally:
        if wb is not None:
            try:
                wb.Close(False)
            except Exception as exc:
                observations["close_error"] = repr(exc)
            del wb
    return observations


def excel_session(action: Callable[[Any, dict[str, Any]], Any]) -> tuple[Any, dict[str, Any]]:
    pythoncom.CoInitialize()
    app = None
    audit: dict[str, Any] = {"started_at": now(), "processes_before": process_snapshot({"excel.exe"})}
    result = None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        hwnd = int(app.Hwnd)
        pid = pid_from_hwnd(hwnd)
        audit.update({"dispatch": "DispatchEx(Excel.Application)", "hwnd": hwnd, "pid": pid, "hwnd_pid_match_at_creation": pid_from_hwnd(hwnd) == pid, "version": str(app.Version), "build": str(app.Build)})
        audit["security_settings"] = configure_office(app)
        result = action(app, audit)
    except Exception as exc:
        audit["error"] = repr(exc)
        audit["traceback"] = traceback.format_exc()
        raise
    finally:
        if app is not None:
            try:
                audit["open_workbook_count_before_quit"] = int(app.Workbooks.Count)
            except Exception as exc:
                audit["open_workbook_count_before_quit_error"] = repr(exc)
            try:
                app.Quit()
                audit["quit_called"] = True
            except Exception as exc:
                audit["quit_called"] = False
                audit["quit_error"] = repr(exc)
            del app
        gc.collect()
        try:
            pythoncom.CoFreeUnusedLibraries()
        except Exception:
            pass
        pythoncom.CoUninitialize()
        gc.collect()
        pid = audit.get("pid")
        exit_wait_started = time.monotonic()
        for _ in range(150):
            if not process_exists(pid):
                break
            time.sleep(0.1)
        audit["own_pid_residual"] = process_exists(pid)
        audit["own_pid_exit_wait_seconds"] = round(time.monotonic() - exit_wait_started, 3)
        audit["exact_pid_terminate_fallback_used"] = False
        before_pids = {item["pid"] for item in audit.get("processes_before", [])}
        ownership_confirmed = bool(pid and audit.get("hwnd_pid_match_at_creation") and pid not in before_pids)
        no_open_workbooks = audit.get("open_workbook_count_before_quit") == 0
        if audit["own_pid_residual"] and ownership_confirmed and no_open_workbooks:
            try:
                proc = psutil.Process(pid)
                if proc.name().lower() == "excel.exe":
                    proc.terminate()
                    proc.wait(timeout=10)
                    audit["exact_pid_terminate_fallback_used"] = True
                    audit["exact_pid_terminated"] = pid
            except (psutil.NoSuchProcess, psutil.TimeoutExpired, psutil.AccessDenied) as exc:
                audit["exact_pid_terminate_error"] = repr(exc)
            audit["own_pid_residual"] = process_exists(pid)
        audit["processes_after"] = process_snapshot({"excel.exe"})
        audit["finished_at"] = now()
    return result, audit


def baseline_command() -> None:
    baseline = workbook_snapshot(SOURCE)

    def observe(app: Any, audit: dict[str, Any]) -> Any:
        return com_observe_workbook(app, SOURCE, "Microsoft Excel COM")

    excel_obs, excel_audit = excel_session(observe)
    baseline["excel_com_inspection"] = excel_obs
    baseline["excel_com_session"] = excel_audit
    baseline["wps_com_inspection"] = {
        "status": "pending_current_isolation_probe",
        "prior_verified_read_only_observation": {
            "used_range": "$A$1:$Q$281",
            "sheet_count": 1,
            "formula_count": 245,
            "version": "12.0",
            "build": "26895",
            "note": "阶段1在既有WPS进程环境中只读打开所得；本次以当前隔离探测结果补充，不能据此声明安全保存。",
        },
    }
    write_json(BASELINE, baseline)
    write_json(LOGS / "source_original_before.json", file_meta(SOURCE))
    print(BASELINE)


def write_simulated_results(ws: Any) -> dict[str, Any]:
    ws.Range("Q1").Value = Q_HEADER
    ws.Range("R1").Value = R_HEADER
    ws.Range("S1").Value = S_HEADER
    target_rows = []
    for row in range(3, 282):
        quote = ws.Cells(row, 13).Value2
        if quote != TARGET_QUOTE:
            continue
        target_rows.append(row)
        if row == SPECIAL_ROW:
            ws.Cells(row, 18).Value = MULTI_STATUS
            ws.Cells(row, 19).Value = MULTI_NOTE
        else:
            ws.Cells(row, 17).Value = round(100.0 + row / 10.0, 1)
            ws.Cells(row, 18).Value = SIM_STATUS
            ws.Cells(row, 19).Value = SIM_NOTE
    return {"target_count": len(target_rows), "target_rows": target_rows, "special_row": SPECIAL_ROW, "distance_rule": "round(100.0 + Excel行号 / 10.0, 1)，静态模拟数值"}


def excel_test_a() -> dict[str, Any]:
    working_copy = WORKING / "excel_test_a_working.xlsx"
    shutil.copy2(SOURCE, working_copy)
    audit: dict[str, Any] = {"test": "A", "source_copy": str(working_copy), "output": str(EXCEL_A), "copy_hash_matches_source": sha256_file(working_copy) == sha256_file(SOURCE)}

    def action(app: Any, session: dict[str, Any]) -> Any:
        wb = None
        ws = None
        try:
            wb = app.Workbooks.Open(str(working_copy), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True, AddToMru=False)
            ws = wb.Worksheets(1)
            write_result = write_simulated_results(ws)
            wb.SaveAs(str(EXCEL_A), FileFormat=XL_MACRO_ENABLED, AddToMru=False)
            return {"write": write_result, "saved_file_format": int(wb.FileFormat), "saved_full_name": str(wb.FullName)}
        finally:
            if wb is not None:
                wb.Close(False)
            if ws is not None:
                del ws
            if wb is not None:
                del wb

    result, session = excel_session(action)
    audit.update({"operation": result, "session": session})

    def verify(app: Any, verify_session: dict[str, Any]) -> Any:
        return com_observe_workbook(app, EXCEL_A, "Microsoft Excel COM validation")

    verification, verify_audit = excel_session(verify)
    audit.update({"reopen_validation": verification, "reopen_session": verify_audit})
    write_json(LOGS / "excel_test_a_audit.json", audit)
    return audit


def excel_test_b() -> dict[str, Any]:
    shutil.copy2(EXCEL_A, EXCEL_B)
    audit: dict[str, Any] = {"test": "B", "source_copy": str(EXCEL_A), "output": str(EXCEL_B), "copy_hash_matches_a": sha256_file(EXCEL_B) == sha256_file(EXCEL_A)}

    def action(app: Any, session: dict[str, Any]) -> Any:
        wb = None
        ws = None
        try:
            wb = app.Workbooks.Open(str(EXCEL_B), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True, AddToMru=False)
            ws = wb.Worksheets(1)
            old = ws.Range("Q3").Value2
            new = round(float(old) + 0.1, 1)
            ws.Range("Q3").Value = new
            wb.Save()
            return {"modified_cell": "Q3", "old_value": old, "new_value": new, "save_method": "Workbook.Save()", "file_format": int(wb.FileFormat)}
        finally:
            if wb is not None:
                wb.Close(False)
            if ws is not None:
                del ws
            if wb is not None:
                del wb

    result, session = excel_session(action)
    audit.update({"operation": result, "session": session})

    def verify(app: Any, verify_session: dict[str, Any]) -> Any:
        return com_observe_workbook(app, EXCEL_B, "Microsoft Excel COM validation")

    verification, verify_audit = excel_session(verify)
    audit.update({"reopen_validation": verification, "reopen_session": verify_audit})
    write_json(LOGS / "excel_test_b_audit.json", audit)
    return audit


def validate_qrs(sheet: dict[str, Any], expected_q3: float | None = None) -> dict[str, Any]:
    rows = sheet["qrs_values_rows_1_281"]
    headers_ok = rows["1"] == [Q_HEADER, R_HEADER, S_HEADER]
    target_rows = []
    non_target_written = []
    target_problems = []
    for row in range(3, 282):
        quote = next((item["value"] for item in sheet["nonempty_cells"] if item["cell"] == f"M{row}"), None)
        qrs = rows[str(row)]
        if quote == TARGET_QUOTE:
            if any(value is not None for value in qrs):
                target_rows.append(row)
            if row == SPECIAL_ROW:
                if qrs != [None, MULTI_STATUS, MULTI_NOTE]:
                    target_problems.append({"row": row, "actual": qrs, "expected": [None, MULTI_STATUS, MULTI_NOTE]})
            else:
                expected_distance = round(100.0 + row / 10.0, 1)
                if row == 3 and expected_q3 is not None:
                    expected_distance = expected_q3
                expected = [expected_distance, SIM_STATUS, SIM_NOTE]
                if qrs != expected:
                    target_problems.append({"row": row, "actual": qrs, "expected": expected})
        elif any(value is not None for value in qrs):
            non_target_written.append({"row": row, "quote_type": quote, "qrs": qrs})
    return {
        "headers_ok": headers_ok,
        "target_rows_with_results": len(target_rows),
        "target_rows": target_rows,
        "target_problems": target_problems,
        "non_target_written": non_target_written,
        "passed": headers_ok and len(target_rows) == 74 and not target_problems and not non_target_written,
    }


def style_differences(source: Path, output: Path) -> list[dict[str, Any]]:
    source_wb = load_workbook(source, read_only=False, data_only=False)
    output_wb = load_workbook(output, read_only=False, data_only=False, keep_vba=output.suffix.lower() == ".xlsm")
    differences: list[dict[str, Any]] = []
    try:
        source_ws = source_wb.worksheets[0]
        output_ws = output_wb.worksheets[0]
        for row in range(1, 282):
            for column in range(1, 17):
                before = style_signature(source_ws.cell(row, column))
                after = style_signature(output_ws.cell(row, column))
                if before != after:
                    differences.append({"cell": source_ws.cell(row, column).coordinate, "baseline": before, "result": after})
    finally:
        source_wb.close()
        output_wb.close()
    return differences


def compare_output(
    output: Path,
    output_name: str,
    expected_q3: float | None = None,
    prior: Path | None = None,
    audit_path: Path | None = None,
    engine_label: str = "Excel",
) -> dict[str, Any]:
    baseline = read_json(BASELINE)
    result = workbook_snapshot(output)
    base_sheet = baseline["python_read_only_inspection"]["sheets"][0]
    out_sheet = result["python_read_only_inspection"]["sheets"][0]
    exact: list[dict[str, Any]] = []
    expected: list[dict[str, Any]] = []
    unexpected: list[dict[str, Any]] = []
    unverifiable: list[dict[str, Any]] = []
    critical: list[str] = []

    def check(label: str, before: Any, after: Any, critical_label: str | None = None) -> None:
        record = {"item": label, "baseline": before, "result": after}
        if before == after:
            exact.append(record)
        else:
            unexpected.append(record)
            if critical_label:
                critical.append(critical_label)

    check("工作表数量", baseline["python_read_only_inspection"]["sheet_count"], result["python_read_only_inspection"]["sheet_count"], "工作表数量变化")
    check("工作表顺序", baseline["python_read_only_inspection"]["sheet_order"], result["python_read_only_inspection"]["sheet_order"], "工作表顺序或名称变化")
    check("工作表可见状态", base_sheet["state"], out_sheet["state"], "工作表可见状态变化")
    check("原A:P值和公式摘要", base_sheet["original_ap_value_formula_digest"], out_sheet["original_ap_value_formula_digest"], "原值或原公式被覆盖/改变")
    if base_sheet["original_ap_style_digest"] == out_sheet["original_ap_style_digest"]:
        exact.append({"item": "原A:P单元格样式摘要", "baseline": base_sheet["original_ap_style_digest"], "result": out_sheet["original_ap_style_digest"]})
    else:
        style_diffs = style_differences(SOURCE, output)
        unexpected.append({
            "item": "原A:P单元格样式发生Office保存规范化",
            "baseline_digest": base_sheet["original_ap_style_digest"],
            "result_digest": out_sheet["original_ap_style_digest"],
            "difference_count": len(style_diffs),
            "differences": style_diffs,
            "assessment": "非目标值和公式未改变；作为非关键但必须披露的样式差异。",
        })
    check("原公式数量", base_sheet["formulas"]["count"], out_sheet["formulas"]["count"], "原公式数量变化")
    check("原公式位置及内容", base_sheet["formulas"]["digest"], out_sheet["formulas"]["digest"], "原公式位置或内容变化")
    check("IF公式数量", base_sheet["formulas"]["if_count"], out_sheet["formulas"]["if_count"], "IF公式数量变化")
    check("DISPIMG公式数量和位置", [base_sheet["formulas"]["dispimg_count"], base_sheet["formulas"]["dispimg_positions"]], [out_sheet["formulas"]["dispimg_count"], out_sheet["formulas"]["dispimg_positions"]], "DISPIMG公式丢失或位置变化")
    check("JPEG/媒体资源哈希", baseline["ooxml_zip_inspection"]["media"], result["ooxml_zip_inspection"]["media"], "JPEG资源丢失或损坏")
    check("cellimages部件", [baseline["ooxml_zip_inspection"]["cell_images_xml_exists"], baseline["ooxml_zip_inspection"]["cell_images_relationship_exists"]], [result["ooxml_zip_inspection"]["cell_images_xml_exists"], result["ooxml_zip_inspection"]["cell_images_relationship_exists"]], "cellimages部件丢失")
    def semantic_cf(value: dict[str, Any]) -> dict[str, Any]:
        normalized = json.loads(json.dumps(value, ensure_ascii=False))
        for entry in normalized.get("entries", []):
            for rule in entry.get("rules", []):
                rule.pop("dxf_id", None)
        return normalized

    check("条件格式语义（规则数/范围/公式/优先级/显示格式）", semantic_cf(base_sheet["conditional_formatting"]), semantic_cf(out_sheet["conditional_formatting"]), "条件格式变化")
    if base_sheet["conditional_formatting"] != out_sheet["conditional_formatting"] and semantic_cf(base_sheet["conditional_formatting"]) == semantic_cf(out_sheet["conditional_formatting"]):
        expected.append({"item": "Excel保存重排条件格式内部dxfId", "baseline": base_sheet["conditional_formatting"], "result": out_sheet["conditional_formatting"], "reason": "规则数、范围、公式、优先级和差异格式语义均一致，仅内部索引重排。"})
    def semantic_filter(value: dict[str, Any]) -> dict[str, Any]:
        columns = []
        for item in value.get("columns", []):
            normalized = {"col_id": item.get("col_id"), "hidden_button": item.get("hidden_button"), "show_button": item.get("show_button")}
            if "filters" in item:
                normalized["equal_values"] = sorted(str(entry) for entry in item.get("filters", []))
                normalized["blank"] = item.get("blank")
            elif item.get("custom_filters") and all(entry.get("operator") in (None, "equal") for entry in item["custom_filters"]):
                normalized["equal_values"] = sorted(str(entry.get("value")) for entry in item["custom_filters"])
                normalized["blank"] = None
            else:
                normalized["raw_condition"] = item
            columns.append(normalized)
        return {"range": value.get("range"), "columns": columns}

    check("AutoFilter范围和筛选语义", semantic_filter(base_sheet["auto_filter"]), semantic_filter(out_sheet["auto_filter"]), "原筛选条件丢失或变化")
    if base_sheet["auto_filter"] != out_sheet["auto_filter"] and semantic_filter(base_sheet["auto_filter"]) == semantic_filter(out_sheet["auto_filter"]):
        expected.append({"item": "Office保存等价改写AutoFilter内部表示", "baseline": base_sheet["auto_filter"], "result": out_sheet["auto_filter"], "reason": "范围、列和等于‘非线路报价’的筛选语义不变；普通Filters被序列化为CustomFilters equal。"})
    check("隐藏行", base_sheet["hidden_rows"], out_sheet["hidden_rows"], "隐藏行异常变化")
    check("隐藏列", base_sheet["hidden_columns"], out_sheet["hidden_columns"], "隐藏列异常变化")
    check("冻结窗格", base_sheet["freeze_panes"], out_sheet["freeze_panes"], "冻结窗格变化")
    check("已定义名称", baseline["python_read_only_inspection"]["defined_names"], result["python_read_only_inspection"]["defined_names"], "已定义名称变化")
    check("合并单元格", base_sheet["merged_cells"], out_sheet["merged_cells"], "合并单元格变化")
    check("外部链接", [baseline["python_read_only_inspection"]["external_links_count"], baseline["ooxml_zip_inspection"]["external_link_parts"], baseline["ooxml_zip_inspection"]["external_relationships"]], [result["python_read_only_inspection"]["external_links_count"], result["ooxml_zip_inspection"]["external_link_parts"], result["ooxml_zip_inspection"]["external_relationships"]])
    check("ActiveX对象", baseline["ooxml_zip_inspection"]["active_x_parts"], result["ooxml_zip_inspection"]["active_x_parts"])
    check("图表部件", baseline["ooxml_zip_inspection"]["chart_parts"], result["ooxml_zip_inspection"]["chart_parts"])
    check("数据验证", base_sheet["data_validations"], out_sheet["data_validations"])

    qrs = validate_qrs(out_sheet, expected_q3=expected_q3)
    expected.append({"item": "文件格式从.xlsx变为真正的宏启用OOXML", "result": {"suffix": output.suffix.lower(), "macro_content_type": result["ooxml_zip_inspection"]["content_type_macro_enabled"], "zip_test": result["ooxml_zip_inspection"]["zip_test"]}})
    expected.append({"item": "新增Q/R/S表头和74行模拟测试结果", "result": qrs})
    if not result["ooxml_zip_inspection"]["content_type_macro_enabled"] or result["ooxml_zip_inspection"]["zip_test"] is not None:
        critical.append("输出不是有效宏启用OOXML")
    if not qrs["passed"]:
        critical.append("Q/R/S目标行或非目标行写入不符合要求")
    if prior is not None:
        prior_snapshot = workbook_snapshot(prior)
        prior_qrs = prior_snapshot["python_read_only_inspection"]["sheets"][0]["qrs_values_rows_1_281"]
        current_qrs = out_sheet["qrs_values_rows_1_281"]
        diffs = [{"row": row, "before": prior_qrs[row], "after": current_qrs[row]} for row in prior_qrs if prior_qrs[row] != current_qrs[row]]
        expected.append({"item": "已有.xlsm再次保存仅修改一个模拟结果单元格", "differences": diffs})
        if diffs != [{"row": "3", "before": [100.3, SIM_STATUS, SIM_NOTE], "after": [100.4, SIM_STATUS, SIM_NOTE]}]:
            critical.append("测试B除Q3外出现额外模拟结果变化")

    source_current = file_meta(SOURCE)
    source_expected = baseline["file"]
    source_unchanged = all(source_current[key] == source_expected[key] for key in ("size_bytes", "modified_time", "sha256"))
    if source_unchanged:
        exact.append({"item": "原样表哈希、大小、修改时间", "baseline": source_expected, "result": source_current})
    else:
        unexpected.append({"item": "原样表哈希、大小、修改时间", "baseline": source_expected, "result": source_current})
        critical.append("原样表发生变化")

    if audit_path is None:
        audit_path = LOGS / ("excel_test_a_audit.json" if output == EXCEL_A else "excel_test_b_audit.json")
    audit = read_json(audit_path) if audit_path.exists() else {}
    baseline_engine_observation = baseline.get("wps_com_inspection", {}).get("isolated_read_only_workbook_observation", {}) if engine_label == "WPS" else baseline.get("excel_com_inspection", {})
    check(
        f"{engine_label} COM已定义名称",
        baseline_engine_observation.get("defined_names"),
        audit.get("reopen_validation", {}).get("defined_names"),
        "已定义名称变化",
    )
    reopen_ok = bool(audit.get("reopen_validation", {}).get("opened")) and audit.get("reopen_validation", {}).get("file_format") == 52
    residual = bool(audit.get("session", {}).get("own_pid_residual")) or bool(audit.get("reopen_session", {}).get("own_pid_residual"))
    if reopen_ok:
        exact.append({"item": f"{engine_label}宏禁用只读重开且FileFormat=52", "result": True})
    else:
        unexpected.append({"item": f"{engine_label}宏禁用只读重开且FileFormat=52", "result": audit.get("reopen_validation")})
        critical.append(f"{engine_label}无法按FileFormat=52重新打开")
    if residual:
        unexpected.append({"item": f"{engine_label}专属进程释放", "result": False})
        critical.append(f"{engine_label}专属进程残留")
    else:
        exact.append({"item": f"{engine_label}专属进程释放", "result": True})

    baseline_com = baseline_engine_observation
    base_text = baseline_com.get("sheets", [{}])[0].get("dispimg_cells")
    out_text = audit.get("reopen_validation", {}).get("sheets", [{}])[0].get("dispimg_cells")
    if base_text and out_text:
        if base_text == out_text:
            exact.append({"item": f"{engine_label}中DISPIMG单元格显示文本/公式/值与源一致", "baseline": base_text, "result": out_text})
        else:
            unverifiable.append({"item": f"{engine_label}中DISPIMG视觉显示", "baseline": base_text, "result": out_text, "reason": "COM文本/值不完全一致，需结合视觉导出；资源与公式另行核对。"})
    else:
        unverifiable.append({"item": "DISPIMG视觉显示位置", "reason": "未取得可比较的COM显示信息。"})

    visual_path = LOGS / "visual_capture_audit.json"
    if visual_path.exists():
        visual = read_json(visual_path)
        label_map = {EXCEL_A: "excel_a", EXCEL_B: "excel_b", WPS_C: "wps_c", WPS_D: "wps_d"}
        visual_label = label_map.get(output)
        visual_entry = visual.get("files", {}).get(visual_label, {}) if visual_label else {}
        captures = visual_entry.get("captures", [])
        manual = visual.get("manual_review", {}).get(visual_label, {}) if visual_label else {}
        visual_record = {
            "item": "WPS渲染下5个DISPIMG位置的视觉检查",
            "captures": [{"cell": item.get("cell"), "output": item.get("output"), "exported": item.get("exported"), "image": item.get("image")} for item in captures],
            "source_comparison": visual.get("source_comparisons", {}).get(visual_label),
            "manual_review": manual,
        }
        if len(captures) == 5 and all(item.get("exported") for item in captures) and manual.get("passed"):
            exact.append(visual_record)
        else:
            unverifiable.append({**visual_record, "reason": "缺少完整截图或尚未记录人工视觉复核结论。"})

    report = {
        "generated_at": now(),
        "baseline": str(BASELINE),
        "result_file": result["file"],
        "summary": {
            "compatible_pass": not critical,
            "critical_failure_count": len(critical),
            "critical_failures": critical,
            "fully_consistent_count": len(exact),
            "expected_change_count": len(expected),
            "unexpected_change_count": len(unexpected),
            "unverifiable_count": len(unverifiable),
        },
        "完全一致": exact,
        "预期变化": expected,
        "非预期变化": unexpected,
        "无法验证": unverifiable,
        "output_snapshot": result,
    }
    write_json(EXPECTED / output_name, report)
    return report


def wps_probe() -> dict[str, Any]:
    pythoncom.CoInitialize()
    app = None
    names = {"et.exe", "wps.exe", "wpsoffice.exe", "wpscloudsvr.exe"}
    report: dict[str, Any] = {"started_at": now(), "processes_before": process_snapshot(names)}
    before_pids = {item["pid"] for item in report["processes_before"]}
    try:
        app = win32com.client.DispatchEx("KET.Application")
        report["dispatch_succeeded"] = True
        report["version"] = json_value(getattr(app, "Version", None))
        report["build"] = json_value(getattr(app, "Build", None))
        hwnd_candidates = {}
        for name in ("Hwnd", "HWND", "MainWindowHandle"):
            try:
                hwnd_candidates[name] = json_value(getattr(app, name))
            except Exception as exc:
                hwnd_candidates[name] = {"error": repr(exc)}
        report["hwnd_candidates"] = hwnd_candidates
        report["security_settings"] = configure_office(app)
        time.sleep(1.0)
        after_dispatch = process_snapshot(names)
        report["processes_after_dispatch"] = after_dispatch
        new_processes = [item for item in after_dispatch if item["pid"] not in before_pids]
        report["new_processes"] = new_processes
        numeric_hwnds = [value for value in hwnd_candidates.values() if isinstance(value, int) and value > 0]
        hwnd_pids = {pid_from_hwnd(hwnd) for hwnd in numeric_hwnds} - {None}
        report["hwnd_pids"] = sorted(hwnd_pids)
        new_et_pids = {item["pid"] for item in new_processes if item["name"] == "et.exe"}
        isolation = len(new_et_pids) == 1 and bool(hwnd_pids & new_et_pids)
        report["isolation_confirmed_before_quit"] = isolation
        if isolation:
            own_pid = next(iter(new_et_pids))
            report["own_pid"] = own_pid
            app.Quit()
            report["quit_called"] = True
            del app
            app = None
            for _ in range(50):
                if not psutil.pid_exists(own_pid):
                    break
                time.sleep(0.1)
            report["own_pid_residual"] = psutil.pid_exists(own_pid)
            final_processes = process_snapshot(names)
            report["processes_final"] = final_processes
            final_pids = {item["pid"] for item in final_processes}
            report["preexisting_processes_unchanged"] = before_pids <= final_pids
            report["safe_for_save_tests"] = not report["own_pid_residual"] and report["preexisting_processes_unchanged"]
        else:
            report["quit_called"] = False
            report["quit_not_called_reason"] = "无法证明COM对象属于新建隔离ET进程；调用Quit可能影响用户既有WPS。"
            report["safe_for_save_tests"] = False
            report["processes_final"] = process_snapshot(names)
    except Exception as exc:
        report["dispatch_succeeded"] = False
        report["error"] = repr(exc)
        report["traceback"] = traceback.format_exc()
        report["safe_for_save_tests"] = False
    finally:
        if app is not None:
            # Safety: do not call Quit when ownership/isolation is not proven.
            del app
        gc.collect()
        pythoncom.CoUninitialize()
        report["finished_at"] = now()
    write_json(LOGS / "wps_isolation_probe.json", report)
    baseline = read_json(BASELINE)
    baseline["wps_com_inspection"]["current_isolation_probe"] = report
    baseline["wps_com_inspection"]["status"] = "isolation_confirmed" if report.get("safe_for_save_tests") else "blocked_unconfirmed_isolation"
    write_json(BASELINE, baseline)
    print(LOGS / "wps_isolation_probe.json")
    return report


def visible_window_fingerprint(processes: list[dict[str, Any]]) -> list[list[Any]]:
    return sorted([
        [item["pid"], window["hwnd"], window["title"]]
        for item in processes
        for window in item.get("windows", [])
        if window.get("visible")
    ])


def wps_session(action: Callable[[Any, dict[str, Any]], Any]) -> tuple[Any, dict[str, Any]]:
    """Run an action only after proving the new KET object owns a new ET PID."""
    pythoncom.CoInitialize()
    app = None
    names = {"et.exe", "wps.exe", "wpsoffice.exe", "wpscloudsvr.exe"}
    before = process_snapshot(names)
    before_pids = {item["pid"] for item in before}
    audit: dict[str, Any] = {"started_at": now(), "processes_before": before, "visible_windows_before": visible_window_fingerprint(before)}
    result = None
    isolated = False
    own_pids: set[int] = set()
    try:
        app = win32com.client.DispatchEx("KET.Application")
        time.sleep(0.8)
        hwnd = int(getattr(app, "Hwnd"))
        pid = pid_from_hwnd(hwnd)
        after_dispatch = process_snapshot(names)
        new_processes = [item for item in after_dispatch if item["pid"] not in before_pids]
        new_et_pids = {item["pid"] for item in new_processes if item["name"] == "et.exe"}
        own_pids = {item["pid"] for item in new_processes}
        isolated = pid is not None and pid in new_et_pids and len(new_et_pids) == 1
        audit.update({
            "dispatch": "DispatchEx(KET.Application)",
            "hwnd": hwnd,
            "pid": pid,
            "version": json_value(getattr(app, "Version", None)),
            "build": json_value(getattr(app, "Build", None)),
            "processes_after_dispatch": after_dispatch,
            "new_processes": new_processes,
            "isolation_confirmed": isolated,
        })
        if not isolated:
            raise RuntimeError("WPS KET实例无法与唯一新建ET进程关联；为保护用户现有WPS，保存动作已阻止。")
        audit["security_settings"] = configure_office(app)
        result = action(app, audit)
    except Exception as exc:
        audit["error"] = repr(exc)
        audit["traceback"] = traceback.format_exc()
    finally:
        if app is not None:
            if isolated:
                try:
                    app.Quit()
                    audit["quit_called"] = True
                except Exception as exc:
                    audit["quit_called"] = False
                    audit["quit_error"] = repr(exc)
            else:
                audit["quit_called"] = False
                audit["quit_not_called_reason"] = "实例所有权未确认，避免影响用户已有WPS。"
            del app
        gc.collect()
        try:
            pythoncom.CoFreeUnusedLibraries()
        except Exception:
            pass
        pythoncom.CoUninitialize()
        gc.collect()
        exit_wait_started = time.monotonic()
        for _ in range(150):
            if not any(psutil.pid_exists(pid) for pid in own_pids):
                break
            time.sleep(0.1)
        final = process_snapshot(names)
        final_pids = {item["pid"] for item in final}
        audit["own_pids"] = sorted(own_pids)
        audit["own_process_residual_pids"] = sorted(pid for pid in own_pids if pid in final_pids)
        audit["own_pid_residual"] = bool(audit["own_process_residual_pids"])
        audit["own_pid_exit_wait_seconds"] = round(time.monotonic() - exit_wait_started, 3)
        audit["processes_after"] = final
        audit["preexisting_processes_unchanged"] = before_pids <= final_pids
        audit["visible_windows_after"] = visible_window_fingerprint(final)
        audit["preexisting_visible_windows_unchanged"] = audit["visible_windows_before"] == audit["visible_windows_after"]
        audit["finished_at"] = now()
    return result, audit


def wps_baseline_observation() -> tuple[dict[str, Any], dict[str, Any]]:
    def observe(app: Any, audit: dict[str, Any]) -> Any:
        return com_observe_workbook(app, SOURCE, "WPS KET COM isolated read-only")

    observation, session = wps_session(observe)
    baseline = read_json(BASELINE)
    baseline["wps_com_inspection"]["isolated_read_only_workbook_observation"] = observation
    baseline["wps_com_inspection"]["isolated_read_only_session"] = session
    write_json(BASELINE, baseline)
    return observation, session


def wps_test_c() -> dict[str, Any]:
    working_copy = WORKING / "wps_test_c_working.xlsx"
    candidate = WORKING / "wps_test_c_candidate.xlsm"
    shutil.copy2(SOURCE, working_copy)
    if candidate.exists():
        candidate.unlink()
    audit: dict[str, Any] = {"test": "C", "source_copy": str(working_copy), "candidate": str(candidate), "output": str(WPS_C), "copy_hash_matches_source": sha256_file(working_copy) == sha256_file(SOURCE)}

    def action(app: Any, session: dict[str, Any]) -> Any:
        wb = None
        ws = None
        try:
            wb = app.Workbooks.Open(str(working_copy), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True, AddToMru=False)
            ws = wb.Worksheets(1)
            write_result = write_simulated_results(ws)
            wb.SaveAs(str(candidate), XL_MACRO_ENABLED)
            return {"write": write_result, "saved_file_format": int(wb.FileFormat), "saved_full_name": str(wb.FullName)}
        finally:
            if wb is not None:
                wb.Close(False)
            if ws is not None:
                del ws
            if wb is not None:
                del wb

    try:
        result, session = wps_session(action)
        audit.update({"operation": result, "session": session})
    except Exception as exc:
        audit.update({"operation_error": repr(exc)})
        # The session audit is unavailable when the action raises before return;
        # candidate evidence is quarantined below if it exists.
    valid = False
    validation_error = None
    if candidate.exists():
        try:
            package = ooxml_snapshot(candidate)
            valid = package["zip_test"] is None and package["content_type_macro_enabled"]
            audit["candidate_ooxml"] = package
        except Exception as exc:
            validation_error = repr(exc)
    audit["candidate_valid_macro_enabled_ooxml"] = valid
    if not valid:
        audit["validation_error"] = validation_error or "ContentType不是macroEnabled或ZIP校验失败"
        if candidate.exists():
            quarantine = WORKING / "WPS_样表转换失败候选_不得业务使用.xlsm"
            if quarantine.exists():
                quarantine.unlink()
            candidate.replace(quarantine)
            audit["quarantined_candidate"] = str(quarantine)
        audit["passed"] = False
        write_json(LOGS / "wps_test_c_audit.json", audit)
        return audit
    if WPS_C.exists():
        WPS_C.unlink()
    candidate.replace(WPS_C)

    def verify_wps(app: Any, verify_session: dict[str, Any]) -> Any:
        return com_observe_workbook(app, WPS_C, "WPS KET COM validation")

    verification, verify_session = wps_session(verify_wps)
    audit.update({"reopen_validation": verification, "reopen_session": verify_session})

    def verify_excel(app: Any, verify_session: dict[str, Any]) -> Any:
        return com_observe_workbook(app, WPS_C, "Microsoft Excel cross-validation")

    excel_verification, excel_session_audit = excel_session(verify_excel)
    audit.update({"excel_cross_validation": excel_verification, "excel_cross_validation_session": excel_session_audit})
    audit["passed"] = bool(verification.get("opened")) and verification.get("file_format") == 52 and not verify_session.get("own_pid_residual")
    write_json(LOGS / "wps_test_c_audit.json", audit)
    return audit


def wps_test_d() -> dict[str, Any]:
    shutil.copy2(WPS_C, WPS_D)
    audit: dict[str, Any] = {"test": "D", "source_copy": str(WPS_C), "output": str(WPS_D), "copy_hash_matches_c": sha256_file(WPS_D) == sha256_file(WPS_C)}

    def action(app: Any, session: dict[str, Any]) -> Any:
        wb = None
        ws = None
        try:
            wb = app.Workbooks.Open(str(WPS_D), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True, AddToMru=False)
            ws = wb.Worksheets(1)
            old = ws.Range("Q3").Value2
            new = round(float(old) + 0.1, 1)
            ws.Range("Q3").Value = new
            wb.Save()
            return {"modified_cell": "Q3", "old_value": old, "new_value": new, "save_method": "Workbook.Save()", "file_format": int(wb.FileFormat)}
        finally:
            if wb is not None:
                wb.Close(False)
            if ws is not None:
                del ws
            if wb is not None:
                del wb

    result, session = wps_session(action)
    audit.update({"operation": result, "session": session})

    def verify(app: Any, verify_session: dict[str, Any]) -> Any:
        return com_observe_workbook(app, WPS_D, "WPS KET COM validation")

    verification, verify_session = wps_session(verify)
    audit.update({"reopen_validation": verification, "reopen_session": verify_session})
    audit["passed"] = bool(verification.get("opened")) and verification.get("file_format") == 52 and not verify_session.get("own_pid_residual")
    write_json(LOGS / "wps_test_d_audit.json", audit)
    return audit


def wps_tests_command() -> None:
    probe = read_json(LOGS / "wps_isolation_probe.json")
    if not probe.get("safe_for_save_tests"):
        print("WPS isolation was not confirmed; save tests skipped.")
        return
    wps_baseline_observation()
    test_c = wps_test_c()
    if not test_c.get("passed") or not WPS_C.exists():
        print("WPS test C failed; test D skipped.")
        return
    comparison_c = compare_output(WPS_C, "WPS_样表转换结果_comparison.json", audit_path=LOGS / "wps_test_c_audit.json", engine_label="WPS")
    if not comparison_c["summary"]["compatible_pass"]:
        print("WPS test C comparison has critical failures; test D skipped.")
        return
    test_d = wps_test_d()
    if test_d.get("passed") and WPS_D.exists():
        compare_output(WPS_D, "WPS_xlsm再次保存结果_comparison.json", expected_q3=100.4, prior=WPS_C, audit_path=LOGS / "wps_test_d_audit.json", engine_label="WPS")
    print(WPS_C)
    if WPS_D.exists():
        print(WPS_D)


def macro_fixture() -> dict[str, Any]:
    fixture = FIXTURES / "office_spike_macro_fixture.xlsm"
    report: dict[str, Any] = {"started_at": now(), "fixture": str(fixture), "trust_center_changed": False}

    def create(app: Any, audit: dict[str, Any]) -> Any:
        wb = None
        ws = None
        try:
            wb = app.Workbooks.Add()
            ws = wb.Worksheets(1)
            ws.Name = "MacroFixture"
            ws.Range("A1").Value = 1
            ws.Range("B1").Value = 2
            ws.Range("C1").Formula = "=SUM(A1:B1)"
            ws.Range("Z1").Value = "NOT_RUN"
            shape = ws.Shapes.AddShape(1, 20, 40, 160, 40)
            shape.TextFrame.Characters().Text = "TASK-003A 安全形状"
            del shape
            try:
                project = wb.VBProject
                standard = project.VBComponents.Add(1)
                standard.Name = "SpikeModule"
                standard.CodeModule.AddFromString('Option Explicit\nPublic Sub SpikeButton()\n    Worksheets("MacroFixture").Range("Z2").Value = "BUTTON_RAN"\nEnd Sub')
                this_wb = project.VBComponents.Item("ThisWorkbook")
                this_wb.CodeModule.AddFromString('Private Sub Workbook_Open()\n    Worksheets("MacroFixture").Range("Z1").Value = "WORKBOOK_OPEN_RAN"\nEnd Sub')
                button = ws.Buttons().Add(20, 100, 160, 30)
                button.Caption = "安全测试按钮（不自动运行）"
                button.OnAction = "SpikeButton"
                del button, this_wb, standard, project
            except Exception as exc:
                return {"created": False, "blocked": True, "blocker": repr(exc), "blocker_stage": "访问Excel VBA工程对象模型", "hresult": getattr(exc, "hresult", None)}
            wb.SaveAs(str(fixture), FileFormat=XL_MACRO_ENABLED, AddToMru=False)
            return {"created": True, "blocked": False, "file_format": int(wb.FileFormat)}
        finally:
            if wb is not None:
                wb.Close(False)
            if ws is not None:
                del ws
            if wb is not None:
                del wb

    result, session = excel_session(create)
    report.update({"creation": result, "creation_session": session})
    if result and result.get("created"):
        before_vba = ooxml_snapshot(fixture)
        before_modules = extract_vba_sources(fixture)

        def verify(app: Any, audit: dict[str, Any]) -> Any:
            wb = None
            try:
                wb = app.Workbooks.Open(str(fixture), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True, AddToMru=False)
                ws = wb.Worksheets("MacroFixture")
                marker = ws.Range("Z1").Value2
                has_vba = bool(wb.HasVBProject)
                formula = json_value(ws.Range("C1").Formula)
                shape_count = int(ws.Shapes.Count)
                button_count = int(ws.Buttons().Count)
                wb.Save()
                return {
                    "marker_after_macro_disabled_open": marker,
                    "workbook_open_did_not_run": marker == "NOT_RUN",
                    "has_vba_project": has_vba,
                    "file_format": int(wb.FileFormat),
                    "formula_c1": formula,
                    "shape_count": shape_count,
                    "button_count": button_count,
                }
            finally:
                if wb is not None:
                    wb.Close(False)

        verification, verify_session = excel_session(verify)
        after_vba = ooxml_snapshot(fixture)
        after_modules = extract_vba_sources(fixture)
        checks = {
            "workbook_open_did_not_run": bool(verification.get("workbook_open_did_not_run")),
            "vba_part_exists_before_and_after": before_vba["vba_project"] and after_vba["vba_project"],
            "vba_components_and_source_preserved": bool(before_modules and before_modules == after_modules),
            "formula_preserved": verification.get("formula_c1") == "=SUM(A1:B1)",
            "shape_and_button_preserved": verification.get("shape_count", 0) >= 2 and verification.get("button_count", 0) >= 1,
            "process_released": not verify_session.get("own_pid_residual"),
        }
        report.update({
            "verification": verification,
            "verification_session": verify_session,
            "vba_part_before_save": before_vba["vba_project"],
            "vba_part_after_save": after_vba["vba_project"],
            "vba_project_sha256_before_save": part_hash(before_vba, "xl/vbaProject.bin"),
            "vba_project_sha256_after_save": part_hash(after_vba, "xl/vbaProject.bin"),
            "vba_modules_before_save": before_modules,
            "vba_modules_after_save": after_modules,
            "checks": checks,
            "passed": all(checks.values()),
        })
    else:
        report["manual_fixture_steps"] = [
            "在Excel中手动新建工作簿并另存为 tests\\fixtures\\office_spike_macro_fixture.xlsm。",
            "在VBA编辑器中添加标准模块 SpikeModule，放入只在手动点击时写入Z2的 SpikeButton 过程。",
            "在 ThisWorkbook 添加 Workbook_Open，仅把 MacroFixture!Z1 从 NOT_RUN 改为 WORKBOOK_OPEN_RAN。",
            "在工作表添加一个窗体按钮并绑定 SpikeButton；再添加普通公式 =SUM(A1:B1) 和一个简单形状。",
            "关闭工作簿；不要降低信任中心设置，不要启用‘信任对VBA工程对象模型的访问’。",
            "把夹具交给本测试脚本，以 AutomationSecurity=3、EnableEvents=False 重新执行宏未运行和VBA保留验证。",
        ]
    write_json(LOGS / "macro_fixture_audit.json", report)
    print(LOGS / "macro_fixture_audit.json")
    return report


def part_hash(package: dict[str, Any], name: str) -> str | None:
    wanted = name.lower()
    for item in package.get("key_parts", []):
        if item.get("name", "").lower() == wanted:
            return item.get("sha256")
    return None


def extract_vba_sources(path: Path) -> list[dict[str, Any]]:
    """Extract VBA source without opening Office or executing macros."""
    from oletools.olevba import VBA_Parser

    parser = VBA_Parser(str(path))
    modules: list[dict[str, Any]] = []
    try:
        if not parser.detect_vba_macros():
            return modules
        for container, stream_path, vba_filename, code in parser.extract_macros():
            if isinstance(code, bytes):
                text = code.decode("utf-8", errors="replace")
            else:
                text = str(code)
            modules.append({
                "container": str(container),
                "stream_path": str(stream_path),
                "vba_filename": str(vba_filename),
                "code": text.replace("\r\n", "\n").replace("\r", "\n").strip(),
            })
    finally:
        parser.close()
    return sorted(modules, key=lambda item: (item["vba_filename"], item["stream_path"]))


def wps_macro_fixture_test() -> dict[str, Any]:
    source_fixture = FIXTURES / "office_spike_macro_fixture.xlsm"
    wps_fixture = FIXTURES / "office_spike_macro_fixture_wps_saved.xlsm"
    report: dict[str, Any] = {"started_at": now(), "source_fixture": str(source_fixture), "wps_saved_fixture": str(wps_fixture)}
    if not source_fixture.exists():
        report.update({"passed": False, "blocked": True, "reason": "Excel宏夹具不存在。"})
        write_json(LOGS / "wps_macro_fixture_audit.json", report)
        return report
    shutil.copy2(source_fixture, wps_fixture)
    before = ooxml_snapshot(wps_fixture)

    def save_action(app: Any, audit: dict[str, Any]) -> Any:
        wb = None
        ws = None
        try:
            wb = app.Workbooks.Open(str(wps_fixture), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True, AddToMru=False)
            ws = wb.Worksheets("MacroFixture")
            observed = {
                "marker_before_save": json_value(ws.Range("Z1").Value2),
                "workbook_open_did_not_run": ws.Range("Z1").Value2 == "NOT_RUN",
                "has_vba_project": bool(wb.HasVBProject),
                "formula_c1": json_value(ws.Range("C1").Formula),
                "shape_count": int(ws.Shapes.Count),
                "button_count": int(ws.Buttons().Count),
                "file_format": int(wb.FileFormat),
            }
            wb.Save()
            return observed
        finally:
            if wb is not None:
                wb.Close(False)
            if ws is not None:
                del ws
            if wb is not None:
                del wb

    save_result, save_session = wps_session(save_action)
    report.update({"wps_save_observation": save_result, "wps_save_session": save_session})

    def wps_verify(app: Any, audit: dict[str, Any]) -> Any:
        wb = None
        ws = None
        try:
            wb = app.Workbooks.Open(str(wps_fixture), UpdateLinks=0, ReadOnly=True, IgnoreReadOnlyRecommended=True, AddToMru=False)
            ws = wb.Worksheets("MacroFixture")
            return {
                "marker_after_reopen": json_value(ws.Range("Z1").Value2),
                "workbook_open_did_not_run": ws.Range("Z1").Value2 == "NOT_RUN",
                "has_vba_project": bool(wb.HasVBProject),
                "formula_c1": json_value(ws.Range("C1").Formula),
                "shape_count": int(ws.Shapes.Count),
                "button_count": int(ws.Buttons().Count),
                "file_format": int(wb.FileFormat),
            }
        finally:
            if wb is not None:
                wb.Close(False)
            if ws is not None:
                del ws
            if wb is not None:
                del wb

    verify_result, verify_session = wps_session(wps_verify)
    report.update({"wps_reopen_observation": verify_result, "wps_reopen_session": verify_session})

    def excel_verify(app: Any, audit: dict[str, Any]) -> Any:
        wb = None
        ws = None
        try:
            wb = app.Workbooks.Open(str(wps_fixture), UpdateLinks=0, ReadOnly=True, IgnoreReadOnlyRecommended=True, AddToMru=False)
            ws = wb.Worksheets("MacroFixture")
            return {
                "marker_after_macro_disabled_open": json_value(ws.Range("Z1").Value2),
                "workbook_open_did_not_run": ws.Range("Z1").Value2 == "NOT_RUN",
                "has_vba_project": bool(wb.HasVBProject),
                "formula_c1": json_value(ws.Range("C1").Formula),
                "shape_count": int(ws.Shapes.Count),
                "button_count": int(ws.Buttons().Count),
                "file_format": int(wb.FileFormat),
            }
        finally:
            if wb is not None:
                wb.Close(False)
            if ws is not None:
                del ws
            if wb is not None:
                del wb

    excel_result, excel_audit = excel_session(excel_verify)
    after = ooxml_snapshot(wps_fixture)
    source_vba_modules = extract_vba_sources(source_fixture)
    wps_saved_vba_modules = extract_vba_sources(wps_fixture)
    report.update({
        "excel_cross_validation": excel_result,
        "excel_cross_validation_session": excel_audit,
        "vba_project_before": {"exists": before["vba_project"], "sha256": part_hash(before, "xl/vbaProject.bin")},
        "vba_project_after": {"exists": after["vba_project"], "sha256": part_hash(after, "xl/vbaProject.bin")},
        "active_x_parts_before": before["active_x_parts"],
        "active_x_parts_after": after["active_x_parts"],
        "drawing_parts_before": before["drawing_parts"],
        "drawing_parts_after": after["drawing_parts"],
        "source_vba_modules": source_vba_modules,
        "wps_saved_vba_modules": wps_saved_vba_modules,
    })
    vba_semantic_preserved = bool(source_vba_modules and source_vba_modules == wps_saved_vba_modules)
    binary_hash_preserved = part_hash(before, "xl/vbaProject.bin") == part_hash(after, "xl/vbaProject.bin")
    checks = {
        "wps_open_macro_not_run": bool(save_result and save_result.get("workbook_open_did_not_run")),
        "wps_reopen_macro_not_run": bool(verify_result and verify_result.get("workbook_open_did_not_run")),
        "excel_cross_open_macro_not_run": bool(excel_result and excel_result.get("workbook_open_did_not_run")),
        "vba_exists": before["vba_project"] and after["vba_project"],
        "vba_components_and_source_preserved": vba_semantic_preserved,
        "shape_and_button_preserved": bool(verify_result and verify_result.get("shape_count", 0) >= 2 and verify_result.get("button_count", 0) >= 1),
        "formula_preserved": bool(verify_result and verify_result.get("formula_c1") == "=SUM(A1:B1)"),
        "wps_processes_released": not save_session.get("own_pid_residual") and not verify_session.get("own_pid_residual"),
        "user_wps_unchanged": bool(save_session.get("preexisting_processes_unchanged")) and bool(save_session.get("preexisting_visible_windows_unchanged")) and bool(verify_session.get("preexisting_processes_unchanged")) and bool(verify_session.get("preexisting_visible_windows_unchanged")),
    }
    report["vba_binary_hash_preserved"] = binary_hash_preserved
    report["vba_binary_hash_change_assessment"] = "二进制哈希变化，但VBA组件名、类型和源码逐行一致；判定为WPS内部重写元数据，不判宏丢失。" if not binary_hash_preserved and vba_semantic_preserved else None
    report["checks"] = checks
    report["passed"] = all(checks.values())
    report["finished_at"] = now()
    write_json(LOGS / "wps_macro_fixture_audit.json", report)
    print(LOGS / "wps_macro_fixture_audit.json")
    return report


def visual_capture_command() -> None:
    from PIL import Image

    def image_meta(path: Path) -> dict[str, Any] | None:
        if not path.exists() or path.stat().st_size <= 0:
            return None
        with Image.open(path) as image:
            rgba = image.convert("RGBA")
            return {
                "size_bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "pixel_size": list(rgba.size),
                "rgba_pixel_sha256": hashlib.sha256(rgba.tobytes()).hexdigest().upper(),
            }

    visual_root = LOGS / "visuals"
    targets = {
        "source": SOURCE,
        "excel_a": EXCEL_A,
        "excel_b": EXCEL_B,
        "wps_c": WPS_C,
        "wps_d": WPS_D,
    }
    report: dict[str, Any] = {"started_at": now(), "engine": "WPS KET COM Range.CopyPicture + temporary in-memory chart export", "files": {}}
    for label, path in targets.items():
        if not path.exists():
            report["files"][label] = {"path": str(path), "status": "missing"}
            continue
        target_dir = visual_root / label
        target_dir.mkdir(parents=True, exist_ok=True)

        def capture(app: Any, audit: dict[str, Any]) -> Any:
            wb = None
            ws = None
            captures = []
            try:
                wb = app.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True, IgnoreReadOnlyRecommended=True, AddToMru=False)
                ws = wb.Worksheets(1)
                for address in ("O38", "O39", "O58", "O78", "O118"):
                    row = int(address[1:])
                    output = target_dir / f"{address}.png"
                    chart_object = None
                    try:
                        ws.Range(f"N{row-1}:P{row+1}").CopyPicture(1, 2)
                        chart_object = ws.ChartObjects().Add(0, 0, 640, 240)
                        chart_object.Chart.Paste()
                        com_return = json_value(chart_object.Chart.Export(str(output)))
                        meta = image_meta(output)
                        captures.append({"cell": address, "output": str(output), "exported": meta is not None, "com_return": com_return, "image": meta})
                    except Exception as exc:
                        captures.append({"cell": address, "output": str(output), "exported": False, "error": repr(exc)})
                    finally:
                        if chart_object is not None:
                            try:
                                chart_object.Delete()
                            except Exception:
                                pass
                            del chart_object
                return captures
            finally:
                if wb is not None:
                    wb.Close(False)
                if ws is not None:
                    del ws
                if wb is not None:
                    del wb

        captures, session = wps_session(capture)
        report["files"][label] = {"path": str(path), "captures": captures, "session": session}
    source_captures = {item["cell"]: item for item in report["files"].get("source", {}).get("captures", [])}
    comparisons = {}
    for label in ("excel_a", "excel_b", "wps_c", "wps_d"):
        items = []
        for item in report["files"].get(label, {}).get("captures", []):
            source_item = source_captures.get(item["cell"])
            source_image = source_item.get("image") if source_item else None
            result_image = item.get("image")
            items.append({
                "cell": item["cell"],
                "source_png_sha256": source_image.get("sha256") if source_image else None,
                "result_png_sha256": result_image.get("sha256") if result_image else None,
                "png_bytes_identical": bool(source_image and result_image and source_image.get("sha256") == result_image.get("sha256")),
                "source_pixel_size": source_image.get("pixel_size") if source_image else None,
                "result_pixel_size": result_image.get("pixel_size") if result_image else None,
                "rgba_pixels_identical": bool(source_image and result_image and source_image.get("rgba_pixel_sha256") == result_image.get("rgba_pixel_sha256")),
            })
        comparisons[label] = items
    report["source_comparisons"] = comparisons
    report["finished_at"] = now()
    write_json(LOGS / "visual_capture_audit.json", report)
    print(LOGS / "visual_capture_audit.json")


def excel_command() -> None:
    def run_worker(job: str) -> dict[str, Any]:
        completed = subprocess.run(
            [sys.executable, str(Path(__file__).resolve()), "excel-worker", "--job", job],
            cwd=ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=180,
            check=False,
        )
        worker_path = LOGS / f"{job}_worker.json"
        if completed.returncode != 0 or not worker_path.exists():
            raise RuntimeError(f"Excel worker {job} failed: rc={completed.returncode}, stdout={completed.stdout}, stderr={completed.stderr}")
        payload = read_json(worker_path)
        session = payload["session"]
        pid = session.get("pid")
        wait_started = time.monotonic()
        for _ in range(100):
            if not process_exists(pid):
                break
            time.sleep(0.1)
        session["own_pid_residual_before_worker_exit"] = session.get("own_pid_residual")
        session["own_pid_residual"] = process_exists(pid)
        session["own_pid_residual_after_worker_exit"] = session["own_pid_residual"]
        session["post_worker_exit_wait_seconds"] = round(time.monotonic() - wait_started, 3)
        payload["worker_process"] = {"returncode": completed.returncode, "stdout": completed.stdout, "stderr": completed.stderr}
        return payload

    a_write = run_worker("excel_a_write")
    a_verify = run_worker("excel_a_verify")
    audit_a = {
        "test": "A",
        "source_copy": str(WORKING / "excel_test_a_working.xlsx"),
        "output": str(EXCEL_A),
        "operation": a_write["result"],
        "session": a_write["session"],
        "reopen_validation": a_verify["result"],
        "reopen_session": a_verify["session"],
        "excel_process_model": "每次COM操作在独立Python辅助进程中执行；Excel在辅助进程退出后释放。",
    }
    write_json(LOGS / "excel_test_a_audit.json", audit_a)
    compare_output(EXCEL_A, "Excel_样表转换结果_comparison.json")

    b_write = run_worker("excel_b_write")
    b_verify = run_worker("excel_b_verify")
    audit_b = {
        "test": "B",
        "source_copy": str(EXCEL_A),
        "output": str(EXCEL_B),
        "operation": b_write["result"],
        "session": b_write["session"],
        "reopen_validation": b_verify["result"],
        "reopen_session": b_verify["session"],
        "excel_process_model": "每次COM操作在独立Python辅助进程中执行；Excel在辅助进程退出后释放。",
    }
    write_json(LOGS / "excel_test_b_audit.json", audit_b)
    compare_output(EXCEL_B, "Excel_xlsm再次保存结果_comparison.json", expected_q3=100.4, prior=EXCEL_A)
    print(EXCEL_A)
    print(EXCEL_B)


def refresh_comparisons_command() -> None:
    if EXCEL_A.exists():
        compare_output(EXCEL_A, "Excel_样表转换结果_comparison.json")
    if EXCEL_B.exists():
        compare_output(EXCEL_B, "Excel_xlsm再次保存结果_comparison.json", expected_q3=100.4, prior=EXCEL_A)
    if WPS_C.exists():
        compare_output(WPS_C, "WPS_样表转换结果_comparison.json", audit_path=LOGS / "wps_test_c_audit.json", engine_label="WPS")
    if WPS_D.exists():
        compare_output(WPS_D, "WPS_xlsm再次保存结果_comparison.json", expected_q3=100.4, prior=WPS_C, audit_path=LOGS / "wps_test_d_audit.json", engine_label="WPS")


def excel_worker(job: str) -> None:
    if job == "excel_a_write":
        working_copy = WORKING / "excel_test_a_working.xlsx"
        shutil.copy2(SOURCE, working_copy)

        def action(app: Any, audit: dict[str, Any]) -> Any:
            wb = None
            ws = None
            try:
                wb = app.Workbooks.Open(str(working_copy), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True, AddToMru=False)
                ws = wb.Worksheets(1)
                result = write_simulated_results(ws)
                wb.SaveAs(str(EXCEL_A), FileFormat=XL_MACRO_ENABLED, AddToMru=False)
                result.update({"saved_file_format": int(wb.FileFormat), "saved_full_name": str(wb.FullName), "copy_hash_matches_source": sha256_file(working_copy) == sha256_file(SOURCE)})
                return result
            finally:
                if wb is not None:
                    wb.Close(False)
                if ws is not None:
                    del ws
                if wb is not None:
                    del wb

    elif job == "excel_a_verify":
        def action(app: Any, audit: dict[str, Any]) -> Any:
            return com_observe_workbook(app, EXCEL_A, "Microsoft Excel COM validation")
    elif job == "excel_b_write":
        shutil.copy2(EXCEL_A, EXCEL_B)

        def action(app: Any, audit: dict[str, Any]) -> Any:
            wb = None
            ws = None
            try:
                wb = app.Workbooks.Open(str(EXCEL_B), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True, AddToMru=False)
                ws = wb.Worksheets(1)
                old = ws.Range("Q3").Value2
                new = round(float(old) + 0.1, 1)
                ws.Range("Q3").Value = new
                wb.Save()
                return {"modified_cell": "Q3", "old_value": old, "new_value": new, "save_method": "Workbook.Save()", "file_format": int(wb.FileFormat)}
            finally:
                if wb is not None:
                    wb.Close(False)
                if ws is not None:
                    del ws
                if wb is not None:
                    del wb
    elif job == "excel_b_verify":
        def action(app: Any, audit: dict[str, Any]) -> Any:
            return com_observe_workbook(app, EXCEL_B, "Microsoft Excel COM validation")
    else:
        raise ValueError(job)
    result, session = excel_session(action)
    write_json(LOGS / f"{job}_worker.json", {"job": job, "result": result, "session": session})


def finalize_baseline_excel_process() -> None:
    baseline = read_json(BASELINE)
    session = baseline["excel_com_session"]
    pid = session.get("pid")
    session["own_pid_residual_before_python_process_exit"] = session.get("own_pid_residual")
    session["own_pid_residual_after_python_process_exit"] = process_exists(pid)
    session["process_boundary_required"] = bool(session.get("own_pid_residual_before_python_process_exit")) and not session["own_pid_residual_after_python_process_exit"]
    write_json(BASELINE, baseline)


def finalize_source_integrity() -> None:
    before = read_json(LOGS / "source_original_before.json")
    after = file_meta(SOURCE)
    result = {"checked_at": now(), "before": before, "after": after, "unchanged": before == after}
    write_json(LOGS / "source_original_after.json", result)
    print(json.dumps(result, ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=("baseline", "baseline-excel-finalize", "excel", "excel-worker", "refresh-comparisons", "wps-probe", "wps-tests", "macro-fixture", "wps-macro-fixture", "visual-capture", "source-integrity"))
    parser.add_argument("--job", choices=("excel_a_write", "excel_a_verify", "excel_b_write", "excel_b_verify"))
    args = parser.parse_args()
    for path in (WORKING, EXPECTED, LOGS, FIXTURES):
        path.mkdir(parents=True, exist_ok=True)
    if args.command == "baseline":
        baseline_command()
    elif args.command == "baseline-excel-finalize":
        finalize_baseline_excel_process()
    elif args.command == "excel":
        excel_command()
    elif args.command == "excel-worker":
        if not args.job:
            raise SystemExit("--job is required for excel-worker")
        excel_worker(args.job)
    elif args.command == "refresh-comparisons":
        refresh_comparisons_command()
    elif args.command == "wps-probe":
        wps_probe()
    elif args.command == "wps-tests":
        wps_tests_command()
    elif args.command == "macro-fixture":
        macro_fixture()
    elif args.command == "wps-macro-fixture":
        wps_macro_fixture_test()
    elif args.command == "visual-capture":
        visual_capture_command()
    elif args.command == "source-integrity":
        finalize_source_integrity()


if __name__ == "__main__":
    main()
