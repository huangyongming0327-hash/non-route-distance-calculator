from __future__ import annotations

import json
import shutil
import sys
import time
from pathlib import Path

import psutil
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.performance import PerformanceLogger
from src.ui.main_window import OutputPermissionWorker
from src.workbook.preview import inspect_workbook_detailed


OFFICE_PROCESS_NAMES = {"excel.exe", "et.exe", "wps.exe", "wpscloudsvr.exe"}


def _office_processes() -> list[dict[str, object]]:
    found = []
    for process in psutil.process_iter(["pid", "name"]):
        try:
            name = str(process.info.get("name") or "").lower()
            if name in OFFICE_PROCESS_NAMES:
                found.append({"pid": process.info["pid"], "name": name})
        except (psutil.AccessDenied, psutil.NoSuchProcess):
            continue
    return found


def _make_workbook(path: Path, data_rows: int, header_row: int) -> None:
    book = Workbook(write_only=True)
    sheet = book.create_sheet("性能测试")
    for row_number in range(1, header_row):
        sheet.append([f"说明{row_number}"])
    sheet.append([
        "始发城市", "始发详细地址", "目的城市", "目的详细地址", "车型", "报价类型"
    ])
    for index in range(data_rows):
        sheet.append([
            "广州", f"测试始发地址{index}", "深圳", f"测试目的地址{index}",
            "4.2米", "非线路报价",
        ])
    book.save(path)


def _measure_detection(path: Path, cache_dir: Path) -> dict[str, object]:
    runs = []
    for _ in range(2):
        stage_timings = []

        def timing(stage: str, elapsed_ms: float, status: str, details: dict) -> None:
            stage_timings.append({
                "stage": stage,
                "elapsed_ms": round(elapsed_ms, 2),
                "status": status,
            })

        started = time.perf_counter()
        outcome = inspect_workbook_detailed(path, cache_dir=cache_dir, timing=timing)
        runs.append({
            "cache_hit": outcome.cache_hit,
            "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
            "header_row": outcome.info.header_row,
            "max_row": outcome.info.max_row,
            "stage_timings": stage_timings,
        })
    return {"file_size": path.stat().st_size, "runs": runs}


def _measure_output_permission(directory: Path, request_id: int) -> dict[str, object]:
    logger = PerformanceLogger(ROOT / "temp" / "task007a_benchmark" / "permission.jsonl")
    worker = OutputPermissionWorker(request_id, str(directory), logger)
    result = []
    worker.finished.connect(lambda *values: result.append(values))
    started = time.perf_counter()
    worker.run()
    return {
        "existing_files": sum(1 for _ in directory.iterdir()),
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "status": result[0][1],
    }


def main() -> int:
    benchmark_root = ROOT / "temp" / "task007a_benchmark"
    expected = (ROOT / "temp" / "task007a_benchmark").resolve()
    if benchmark_root.resolve() != expected:
        raise RuntimeError("性能测试临时目录不符合预期。")
    if benchmark_root.exists():
        shutil.rmtree(benchmark_root)
    benchmark_root.mkdir(parents=True)

    before_office = _office_processes()
    sample = next((ROOT / "samples" / "input").glob("*.xlsx"))
    synthetic_2k = benchmark_root / "中文路径" / "两千行 样表.xlsx"
    synthetic_10k = benchmark_root / "带 空格 路径" / "一万行 样表.xlsx"
    synthetic_2k.parent.mkdir(parents=True)
    synthetic_10k.parent.mkdir(parents=True)
    _make_workbook(synthetic_2k, 2_000, 4)
    _make_workbook(synthetic_10k, 10_000, 15)

    empty_output = benchmark_root / "empty_output"
    large_output = benchmark_root / "large_output"
    empty_output.mkdir()
    large_output.mkdir()
    for index in range(1_000):
        (large_output / f"existing_{index}.txt").touch()

    evidence = {
        "measured_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "environment": {
            "python": sys.version.split()[0],
            "platform": sys.platform,
            "real_amap_calls": 0,
        },
        "workbooks": {
            "current_281_rows": _measure_detection(sample, benchmark_root / "cache_281"),
            "synthetic_2000_rows_header_4": _measure_detection(
                synthetic_2k, benchmark_root / "cache_2k"
            ),
            "synthetic_10000_rows_header_15": _measure_detection(
                synthetic_10k, benchmark_root / "cache_10k"
            ),
        },
        "output_directories": {
            "empty": _measure_output_permission(empty_output, 1),
            "contains_1000_files": _measure_output_permission(large_output, 2),
        },
        "office_processes_before": before_office,
        "office_processes_after": _office_processes(),
    }
    evidence["office_processes_started_by_detection"] = [
        item for item in evidence["office_processes_after"]
        if item not in evidence["office_processes_before"]
    ]
    target = ROOT / "docs" / "evidence" / "task007a_performance.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"EVIDENCE={target}")
    for name, result in evidence["workbooks"].items():
        print(
            f"{name}: first={result['runs'][0]['elapsed_ms']}ms, "
            f"cached={result['runs'][1]['elapsed_ms']}ms"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
