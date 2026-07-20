from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
import json
from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.application.runner import PrototypeRunner
from src.domain.models import OfficeEngine, TaskMode, WorkbookSelection
from src.utils.files import fingerprint
from src.workbook.preview import inspect_workbook
from src.workbook.verification import validate_prototype_output, write_validation_report


OUTPUT_DIR = ROOT / "samples" / "expected" / "address_confirmation"
EVIDENCE_DIR = ROOT / "docs" / "evidence"
OUTPUT_NAMES = {
    OfficeEngine.EXCEL: "Excel_样表_可信地址库结果_TASK-005A_待验收.xlsm",
    OfficeEngine.WPS: "WPS_样表_可信地址库结果_TASK-005A_待验收.xlsm",
}


def _status_counts(path: Path, sheet_name: str, status_column: int) -> dict[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=True, keep_links=False)
    try:
        sheet = workbook[sheet_name]
        return dict(Counter(
            str(sheet.cell(row, status_column).value)
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row, status_column).value
        ))
    finally:
        workbook.close()


def run_engine(engine: OfficeEngine) -> dict:
    source = next((ROOT / "samples" / "input").glob("*.xlsx"))
    before = fingerprint(source)
    info = inspect_workbook(source)
    selection = WorkbookSelection(
        str(source), info.selected_sheet, info.header_row, info.recommended_mapping
    )
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = OUTPUT_DIR / OUTPUT_NAMES[engine]
    task_id = f"task005a_{engine.value}_{datetime.now().strftime('%Y%m%d%H%M%S')}"
    summary = PrototypeRunner(ROOT).run(
        selection,
        engine=engine,
        mode=TaskMode.DRIVING_REAL,
        output_dir=OUTPUT_DIR,
        final_path=output,
        task_id=task_id,
        force_all=True,
        force_route_refresh=False,
        checkpoint_batch=10,
        checkpoint_seconds=60.0,
    )
    validation = validate_prototype_output(
        source,
        output,
        sheet_name=info.selected_sheet,
        header_row=info.header_row,
        fields=info.recommended_mapping,
        results=info.result_columns,
        expected_source_fingerprint=before,
        mode=TaskMode.DRIVING_REAL,
    )
    status_counts = _status_counts(output, info.selected_sheet, info.result_columns.status)
    task005a_checks = {
        "source_hash_unchanged": fingerprint(source) == before,
        "distance_count_67": validation["distance_count"] == 67,
        "multi_destination_row_248": validation["multi_destination_rows"] == [248],
        "city_conflict_rows_39_78_272": validation["conflict_rows"] == [39, 78, 272],
        "blocked_rows_written": status_counts.get("地址风险过高—未计算", 0) == 6,
        "review_warning_rows_written": status_counts.get("查询成功—定位待复核", 0) == 52,
        "no_old_low_precision_status": "查询成功—定位精度较低" not in status_counts,
        "no_http_needed_after_address_migration": summary.actual_http_calls == 0,
    }
    validation.update({
        "run_summary": summary.__dict__,
        "status_counts": status_counts,
        "task005a_checks": task005a_checks,
        "task005a_failed_checks": [name for name, passed in task005a_checks.items() if not passed],
    })
    validation["passed"] = validation["passed"] and all(task005a_checks.values())
    report_path = EVIDENCE_DIR / f"TASK005A_{engine.value.upper()}_VALIDATION.json"
    write_validation_report(report_path, validation)
    if not validation["passed"]:
        raise RuntimeError(
            f"{engine.value} TASK-005A 验证失败："
            f"{validation['failed_checks']} / {validation['task005a_failed_checks']}"
        )
    return {
        "engine": engine.value,
        "output": str(output),
        "validation": str(report_path),
        "summary": summary.__dict__,
        "status_counts": status_counts,
        "passed": True,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("engine", choices=("excel", "wps", "all"))
    args = parser.parse_args()
    engines = (
        (OfficeEngine.EXCEL, OfficeEngine.WPS)
        if args.engine == "all" else (OfficeEngine(args.engine),)
    )
    results = {engine.value: run_engine(engine) for engine in engines}
    print(json.dumps(results, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
