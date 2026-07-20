from __future__ import annotations

from collections import Counter
from datetime import datetime
import json
import os
from pathlib import Path
import sys
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.security.key_store import SecureKeyStore
from src.utils.files import sha256_file
from src.workbook.preview import RESULT_HEADERS, inspect_workbook


EVIDENCE = ROOT / "docs" / "evidence"
OUTPUT = ROOT / "samples" / "expected" / "driving_real"


def _read(name: str) -> dict[str, Any]:
    return json.loads((EVIDENCE / name).read_text(encoding="utf-8"))


def _output_audit(path: Path) -> dict[str, Any]:
    book = load_workbook(path, read_only=True, data_only=False, keep_vba=True, keep_links=False)
    try:
        ws = book[book.sheetnames[0]]
        headers = [ws.cell(1, col).value for col in range(17, 20)]
        statuses = [str(ws.cell(row, 18).value or "") for row in range(2, ws.max_row + 1)]
        forbidden = ("货车查询成功", "货车可通行", "货车路线", "专业货车已开通", "车型未识别")
        text = "\n".join(
            f"{ws.cell(row, 18).value or ''}\n{ws.cell(row, 19).value or ''}"
            for row in range(2, ws.max_row + 1)
        )
        return {
            "path": str(path),
            "sha256": sha256_file(path),
            "size": path.stat().st_size,
            "headers": headers,
            "headers_ok": headers == list(RESULT_HEADERS),
            "status_counts": dict(Counter(item for item in statuses if item)),
            "forbidden_claims_absent": not any(token in text for token in forbidden),
        }
    finally:
        book.close()


def _secret_scan() -> dict[str, Any]:
    secret = SecureKeyStore(ROOT).get_key()
    if not secret:
        return {"key_loaded": False, "plaintext_findings": []}
    needle = secret.encode("utf-8")
    findings: list[str] = []
    excluded = {".git", ".venv", "cache", "config", "build", "release", "temp"}
    for path in ROOT.rglob("*"):
        if not path.is_file() or any(part in excluded for part in path.relative_to(ROOT).parts):
            continue
        if path.stat().st_size > 12 * 1024 * 1024:
            continue
        try:
            if needle in path.read_bytes():
                findings.append(str(path))
        except OSError:
            continue
    return {"key_loaded": True, "plaintext_findings": findings}


def main() -> None:
    connection = _read("TASK004R_CONNECTION.json")
    small = _read("TASK004R_SMALL_ROUTES.json")
    excel_validation = _read("TASK004R_EXCEL_VALIDATION.json")
    wps_validation = _read("TASK004R_WPS_VALIDATION.json")
    source = next((ROOT / "samples" / "input").glob("*.xlsx"))
    excel = _output_audit(OUTPUT / "Excel_样表_普通驾车距离结果_待验收.xlsm")
    wps = _output_audit(OUTPUT / "WPS_样表_普通驾车距离结果_待验收.xlsm")
    initial_excel_calls = 0
    for path in (ROOT / "logs" / "driving_real").glob("task004r_excel_*/run_summary.json"):
        summary = json.loads(path.read_text(encoding="utf-8")).get("summary", {})
        initial_excel_calls = max(initial_excel_calls, int(summary.get("actual_http_calls", 0)))
    secret_scan = _secret_scan()
    source_text = "\n".join(
        path.read_text(encoding="utf-8", errors="ignore")
        for path in (ROOT / "src").rglob("*.py")
    )
    checks = {
        "connection_passed": connection.get("passed") is True and connection.get("actual_http_calls") == 2,
        "small_routes_passed": small.get("passed") is True and len(small.get("routes", [])) == 3,
        "excel_validation_passed": excel_validation.get("passed") is True,
        "wps_validation_passed": wps_validation.get("passed") is True,
        "source_hash_unchanged": sha256_file(source) == "52011587F8F12A87749087ABEA0ADCE2557874D1B268A55F348E86F8DFA7A951",
        "excel_headers_and_claims": excel["headers_ok"] and excel["forbidden_claims_absent"],
        "wps_headers_and_claims": wps["headers_ok"] and wps["forbidden_claims_absent"],
        "no_truck_endpoint_in_source": "/v4/direction/truck" not in source_text,
        "no_plaintext_key_outside_secure_storage": not secret_scan["plaintext_findings"],
    }
    report = {
        "audited_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source": {"path": str(source), "sha256": sha256_file(source)},
        "http_calls": {
            "connection": connection["actual_http_calls"],
            "small_routes": small["actual_http_calls"],
            "initial_excel_batch": initial_excel_calls,
            "total_before_final_cached_refresh": connection["actual_http_calls"] + small["actual_http_calls"] + initial_excel_calls,
            "final_excel": excel_validation.get("run_summary", {}).get("actual_http_calls"),
            "final_wps": wps_validation.get("run_summary", {}).get("actual_http_calls"),
        },
        "excel": excel,
        "wps": wps,
        "secret_scan": secret_scan,
        "checks": checks,
        "passed": all(checks.values()),
    }
    target = EVIDENCE / "TASK004R_FINAL_AUDIT.json"
    temporary = target.with_name(f".{target.name}.tmp")
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary, target)
    print(json.dumps({"passed": report["passed"], "http_calls": report["http_calls"], "checks": checks}, ensure_ascii=False, indent=2))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
