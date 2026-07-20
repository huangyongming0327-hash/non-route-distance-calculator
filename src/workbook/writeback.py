from __future__ import annotations

from typing import Any

from openpyxl.utils import get_column_letter

from src.domain.models import FieldMapping, ResultColumns, RowOutcome
from src.workbook.preview import RESULT_HEADERS


WARNING_MARKER_NAME = "_TASK004R_WARNING_MARKER"
WARNING_MARKER_VALUE = "TASK004R_DRIVING_V1"
XL_EXPRESSION = 2


def _rgb(red: int, green: int, blue: int) -> int:
    return red + green * 256 + blue * 65536


def _defined_name_exists(workbook: Any, name: str) -> bool:
    for index in range(1, int(workbook.Names.Count) + 1):
        item = workbook.Names(index)
        try:
            current = str(item.Name).split("!")[-1].strip("'")
            if current == name:
                return True
        finally:
            del item
    return False


def _add_warning_rule(ws: Any, address: str, formula: str, *, yellow: bool = False) -> None:
    # Excel/WPS 的 late-bound COM 对命名可选参数支持不一致；完整位置参数
    # (Type, Operator, Formula1, Formula2) 在两端更稳定。
    condition = ws.Range(address).FormatConditions.Add(XL_EXPRESSION, None, formula, None)
    if yellow:
        condition.Interior.Color = _rgb(255, 235, 156)
        condition.Font.Color = _rgb(156, 101, 0)
    else:
        condition.Interior.Color = _rgb(255, 199, 206)
        condition.Font.Color = _rgb(156, 0, 6)
    del condition


def add_tool_conditional_formatting(
    workbook: Any,
    ws: Any,
    fields: FieldMapping,
    results: ResultColumns,
    first_data_row: int,
    last_data_row: int,
) -> bool:
    if _defined_name_exists(workbook, WARNING_MARKER_NAME):
        return False
    workbook.Names.Add(Name=WARNING_MARKER_NAME, RefersTo=f'="{WARNING_MARKER_VALUE}"', Visible=False)
    status_col = get_column_letter(results.status)
    explanation_col = get_column_letter(results.explanation)
    marker = f'{WARNING_MARKER_NAME}="{WARNING_MARKER_VALUE}"'
    # WPS 在给整列区域添加条件格式时会错误平移相对行引用到工作表末端。
    # 用 ROW() + 有界绝对 INDEX 取得当前行，避免相对引用且不使用整列公式。
    status_ref = f"INDEX(${status_col}$1:${status_col}${last_data_row},ROW())&\"\""
    explanation_ref = f"INDEX(${explanation_col}$1:${explanation_col}${last_data_row},ROW())&\"\""
    origin_formula = (
        f'=AND({marker},OR(ISNUMBER(SEARCH("发货城市冲突",{explanation_ref})),'
        f'ISNUMBER(SEARCH("发货地址为空",{explanation_ref}))))'
    )
    destination_formula = (
        f'=AND({marker},OR(ISNUMBER(SEARCH("多目的地",{status_ref})),'
        f'ISNUMBER(SEARCH("目的城市冲突",{explanation_ref})),ISNUMBER(SEARCH("目的地址为空",{explanation_ref}))))'
    )
    warning_formula = (
        f'=AND({marker},OR(ISNUMBER(SEARCH("待确认",{status_ref})),ISNUMBER(SEARCH("地址为空",{status_ref})),'
        f'ISNUMBER(SEARCH("定位失败",{status_ref})),ISNUMBER(SEARCH("解析失败",{status_ref})),'
        f'ISNUMBER(SEARCH("风险过高",{status_ref})),ISNUMBER(SEARCH("地址无效",{status_ref}))))'
    )
    review_formula = (
        f'=AND({marker},OR(ISNUMBER(SEARCH("精度较低",{status_ref})),'
        f'ISNUMBER(SEARCH("定位待复核",{status_ref})),ISNUMBER(SEARCH("人工复核",{status_ref}))))'
    )
    for col in (fields.origin_city, fields.origin_address):
        letter = get_column_letter(col)
        _add_warning_rule(ws, f"{letter}{first_data_row}:{letter}{last_data_row}", origin_formula)
    for col in (fields.destination_city, fields.destination_address):
        letter = get_column_letter(col)
        _add_warning_rule(ws, f"{letter}{first_data_row}:{letter}{last_data_row}", destination_formula)
    warning_range = f"{status_col}{first_data_row}:{explanation_col}{last_data_row}"
    _add_warning_rule(ws, warning_range, warning_formula)
    _add_warning_rule(ws, warning_range, review_formula, yellow=True)
    return True


def write_results(
    workbook: Any,
    ws: Any,
    header_row: int,
    fields: FieldMapping,
    results: ResultColumns,
    outcomes: list[RowOutcome],
    last_data_row: int,
) -> dict[str, Any]:
    for col, header in zip(results.as_dict().values(), RESULT_HEADERS, strict=True):
        ws.Cells(header_row, col).Value2 = header
        ws.Cells(header_row, col).Font.Bold = True
        ws.Cells(header_row, col).WrapText = True
    current_header_height = float(ws.Rows(header_row).RowHeight or 0)
    ws.Rows(header_row).RowHeight = max(current_header_height, 42)
    ws.Columns(results.distance).ColumnWidth = 26
    ws.Columns(results.status).ColumnWidth = 28
    ws.Columns(results.explanation).ColumnWidth = 68
    for outcome in outcomes:
        distance_cell = ws.Cells(outcome.excel_row, results.distance)
        if outcome.distance_km is None:
            distance_cell.ClearContents()
        else:
            distance_cell.Value2 = float(outcome.distance_km)
        distance_cell.NumberFormat = "0.0"
        ws.Cells(outcome.excel_row, results.status).Value2 = outcome.status
        explanation_cell = ws.Cells(outcome.excel_row, results.explanation)
        explanation_cell.Value2 = outcome.explanation
        explanation_cell.WrapText = True
        del distance_cell, explanation_cell
    added = add_tool_conditional_formatting(
        workbook, ws, fields, results, header_row + 1, last_data_row
    )
    return {"written_rows": len(outcomes), "conditional_formatting_added": added}
