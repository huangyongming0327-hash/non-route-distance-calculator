from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
import time
from dataclasses import asdict, dataclass
from pathlib import Path
from threading import Event
from typing import Any, Callable
from xml.etree.ElementTree import iterparse

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.domain.models import FieldMapping, ResultColumns, RowInput
from src.utils.files import reject_xlm_macro_sheets
from src.utils.text import normalized_text


RESULT_HEADERS = (
    "高德普通驾车公路距离（公里）",
    "距离查询状态",
    "距离查询说明",
)

HEADER_KEYWORDS = {
    "quote_type": ("报价类型", "报价"),
    "origin_city": ("始发城市", "发货城市", "始发仓", "发货地"),
    "origin_address": ("始发详细地址", "发货详细地址", "始发地址"),
    "destination_city": ("目的城市", "到达城市", "到达仓", "目的地"),
    "destination_address": ("目的详细地址", "到达详细地址", "目的地址"),
    "vehicle": ("车型", "车长"),
}

HEADER_SCAN_LIMIT = 20
SAMPLE_DATA_ROW_LIMIT = 30
PREVIEW_DATA_ROW_LIMIT = 10
DETECTION_ALGORITHM_VERSION = "task-007a-ooxml-v1"
DEFAULT_CITY_CATALOG = (
    "北京", "上海", "天津", "重庆", "广州", "深圳", "杭州", "佛山", "贵阳", "无锡"
)

ProgressCallback = Callable[[str], None]
TimingCallback = Callable[[str, float, str, dict[str, Any]], None]


class InspectionCancelled(RuntimeError):
    pass


@dataclass(frozen=True)
class WorkbookInfo:
    path: str
    sheets: tuple[str, ...]
    selected_sheet: str
    header_row: int
    header_score: int
    max_row: int
    max_column: int
    actual_last_column: int
    descriptors: tuple[str, ...]
    recommended_mapping: FieldMapping
    result_columns: ResultColumns
    city_catalog: tuple[str, ...]
    preview_rows: tuple[tuple[Any, ...], ...]


@dataclass(frozen=True)
class InspectionOutcome:
    info: WorkbookInfo
    cache_hit: bool


def _check_cancel(cancel_event: Event | None) -> None:
    if cancel_event is not None and cancel_event.is_set():
        raise InspectionCancelled("用户已取消当前检测。")


def _timed(
    stage: str,
    action: Callable[[], Any],
    timing: TimingCallback | None,
    *,
    details: dict[str, Any] | None = None,
) -> Any:
    started = time.perf_counter()
    try:
        value = action()
    except Exception:
        if timing:
            timing(stage, (time.perf_counter() - started) * 1000, "failed", details or {})
        raise
    if timing:
        timing(stage, (time.perf_counter() - started) * 1000, "completed", details or {})
    return value


def _score_values(values: tuple[Any, ...]) -> int:
    score = 0
    for value in values[:80]:
        text = normalized_text(value, remove_all_space=True)
        for keywords in HEADER_KEYWORDS.values():
            if any(keyword in text for keyword in keywords):
                score += 3 if text in keywords else 1
    return score


def _score_header_row(ws: Any, row: int) -> int:
    values = next(
        ws.iter_rows(min_row=row, max_row=row, max_col=min(ws.max_column, 80), values_only=True),
        (),
    )
    return _score_values(tuple(values))


def _detect_header_from_rows(rows: list[tuple[Any, ...]], limit: int = HEADER_SCAN_LIMIT) -> tuple[int, int]:
    candidates = [
        (row_number, _score_values(values))
        for row_number, values in enumerate(rows[:limit], start=1)
    ]
    if not candidates:
        raise ValueError("工作表为空，无法检测表头。")
    row, score = max(candidates, key=lambda item: (item[1], -item[0]))
    if score <= 0:
        raise ValueError("前 20 行未检测到可信表头，请手动选择表头行。")
    return row, score


def detect_header_row(ws: Any, limit: int = HEADER_SCAN_LIMIT) -> tuple[int, int]:
    rows = [
        tuple(row)
        for row in ws.iter_rows(
            min_row=1,
            max_row=min(limit, ws.max_row),
            max_col=min(ws.max_column, 80),
            values_only=True,
        )
    ]
    return _detect_header_from_rows(rows, limit)


def _sample_values_from_rows(
    rows: list[tuple[Any, ...]], header_row: int, col: int, limit: int = 3
) -> list[str]:
    values: list[str] = []
    for values_row in rows[header_row : header_row + SAMPLE_DATA_ROW_LIMIT]:
        raw = values_row[col - 1] if col <= len(values_row) else None
        value = normalized_text(raw)
        if value and value not in values:
            values.append(value)
        if len(values) >= limit:
            break
    return values


def _descriptor_from_rows(rows: list[tuple[Any, ...]], header_row: int, col: int) -> str:
    header_values = rows[header_row - 1] if header_row <= len(rows) else ()
    header = normalized_text(header_values[col - 1] if col <= len(header_values) else None) or "（空表头）"
    samples = " / ".join(_sample_values_from_rows(rows, header_row, col)) or "（无样例）"
    if len(samples) > 60:
        samples = samples[:57] + "..."
    return f"{get_column_letter(col)}｜{header}｜{samples}"


def _address_score(samples: list[str]) -> int:
    markers = ("省", "市", "区", "县", "路", "街", "道", "号", "园", "仓", "镇")
    return sum(sum(marker in sample for marker in markers) for sample in samples)


def _recommend_mapping_from_rows(
    rows: list[tuple[Any, ...]], header_row: int, max_column: int
) -> FieldMapping:
    header_values = rows[header_row - 1] if header_row <= len(rows) else ()
    headers = {
        col: normalized_text(header_values[col - 1] if col <= len(header_values) else None, remove_all_space=True)
        for col in range(1, max_column + 1)
    }

    def best_header(field: str) -> int | None:
        keywords = HEADER_KEYWORDS[field]
        scored = []
        for col, text in headers.items():
            score = max((100 if text == keyword else 60 if keyword in text else 0) for keyword in keywords)
            if score:
                scored.append((score, -col, col))
        return max(scored)[2] if scored else None

    quote = best_header("quote_type")
    vehicle = best_header("vehicle")
    origin_city = best_header("origin_city")
    destination_city = best_header("destination_city")
    if not all((quote, origin_city, destination_city)):
        raise ValueError("无法完整推荐报价、始发或目的列，请在界面手动确认。")

    def adjacent_address(city_col: int) -> int:
        candidates = [col for col in (city_col + 1, city_col - 1) if 1 <= col <= max_column]
        scored = []
        for col in candidates:
            samples = _sample_values_from_rows(rows, header_row, col, limit=8)
            keywords = HEADER_KEYWORDS["origin_address"] + HEADER_KEYWORDS["destination_address"]
            explicit = 50 if any(keyword in headers[col] for keyword in keywords) else 0
            scored.append((explicit + _address_score(samples), -abs(col - city_col), col))
        return max(scored)[2]

    origin_address = best_header("origin_address") or adjacent_address(origin_city)
    destination_address = best_header("destination_address") or adjacent_address(destination_city)
    return FieldMapping(
        quote, origin_city, origin_address, destination_city, destination_address, vehicle
    )


def recommend_mapping(ws: Any, header_row: int) -> FieldMapping:
    max_row = min(ws.max_row, header_row + SAMPLE_DATA_ROW_LIMIT)
    rows = [
        tuple(row)
        for row in ws.iter_rows(
            min_row=1, max_row=max_row, max_col=ws.max_column, values_only=True
        )
    ]
    return _recommend_mapping_from_rows(rows, header_row, ws.max_column)


def actual_last_column(ws: Any) -> int:
    last = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                last = max(last, cell.column)
    return last


def _actual_dimensions_from_ooxml(
    book: Any, ws: Any, cancel_event: Event | None
) -> tuple[int, int]:
    """只解析单个工作表 XML 的单元格引用，不读取图片、图表、宏或样式。"""
    last = 0
    last_row = 0
    cell_count = 0
    with book._archive.open(ws._worksheet_path) as source:  # openpyxl 的只读 OOXML 归档
        for _, element in iterparse(source, events=("end",)):
            if not element.tag.endswith("}c"):
                if element.tag.endswith("}row"):
                    element.clear()
                continue
            cell_count += 1
            if cell_count % 2048 == 0:
                _check_cancel(cancel_event)
            # 样式占位单元格没有 v/f/is；与原 actual_last_column 的非空语义一致。
            if not any(
                child.tag.endswith("}v")
                or child.tag.endswith("}f")
                or child.tag.endswith("}is")
                for child in element
            ):
                element.clear()
                continue
            reference = element.attrib.get("r", "")
            match = re.match(r"([A-Z]+)", reference)
            if match:
                column = 0
                for char in match.group(1):
                    column = column * 26 + ord(char) - 64
                last = max(last, column)
                row_match = re.search(r"(\d+)$", reference)
                if row_match:
                    last_row = max(last_row, int(row_match.group(1)))
            element.clear()
    return last, last_row


def _plan_result_columns_from_header(
    header_values: tuple[Any, ...], max_column: int, last_column: int
) -> ResultColumns:
    expected = tuple(normalized_text(header, remove_all_space=True) for header in RESULT_HEADERS)
    for start in range(1, max(1, max_column - 1)):
        actual = tuple(
            normalized_text(
                header_values[start + offset - 1]
                if start + offset <= len(header_values)
                else None,
                remove_all_space=True,
            )
            for offset in range(3)
        )
        if actual == expected:
            return ResultColumns(start, start + 1, start + 2)
    start = last_column + 1
    return ResultColumns(start, start + 1, start + 2)


def plan_result_columns(ws: Any, header_row: int) -> ResultColumns:
    expected = tuple(normalized_text(header, remove_all_space=True) for header in RESULT_HEADERS)
    for start in range(1, max(1, ws.max_column - 1)):
        actual = tuple(
            normalized_text(ws.cell(header_row, start + offset).value, remove_all_space=True)
            for offset in range(3)
        )
        if actual == expected:
            return ResultColumns(start, start + 1, start + 2)
    start = actual_last_column(ws) + 1
    return ResultColumns(start, start + 1, start + 2)


def _collect_city_catalog(
    ws: Any,
    mapping: FieldMapping,
    header_row: int,
    cancel_event: Event | None,
    max_row: int | None = None,
) -> tuple[str, ...]:
    city_values: set[str] = set(DEFAULT_CITY_CATALOG)
    first_col = min(mapping.origin_city, mapping.destination_city)
    last_col = max(mapping.origin_city, mapping.destination_city)
    origin_offset = mapping.origin_city - first_col
    destination_offset = mapping.destination_city - first_col
    rows = ws.iter_rows(
        min_row=header_row + 1,
        max_row=max_row or ws.max_row,
        min_col=first_col,
        max_col=last_col,
        values_only=True,
    )
    for index, values in enumerate(rows, start=1):
        if index % 512 == 0:
            _check_cancel(cancel_event)
        for offset in (origin_offset, destination_offset):
            value = normalized_text(values[offset], remove_all_space=True)
            if value and len(value) <= 20:
                city_values.add(value.removesuffix("市"))
    return tuple(sorted(city_values))


def _file_key(path: Path, sheet_name: str, requested_header: int | None) -> dict[str, Any]:
    stat = path.stat()
    return {
        "canonical_path": os.path.normcase(str(path.resolve())),
        "file_size": stat.st_size,
        "modified_time_ns": stat.st_mtime_ns,
        "sheet_name": sheet_name,
        "requested_header": requested_header,
        "algorithm_version": DETECTION_ALGORITHM_VERSION,
    }


def _cache_path(cache_dir: Path, key: dict[str, Any]) -> Path:
    digest = hashlib.sha256(
        json.dumps(key, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()
    return cache_dir / f"{digest}.json"


def _redact_sample(value: Any) -> Any:
    if value is None or value == "":
        return value
    text = str(value)
    if len(text) == 1:
        return "*"
    if len(text) == 2:
        return text[0] + "*"
    return text[:2] + "***" + text[-1:]


def _redact_descriptor(descriptor: str) -> str:
    parts = descriptor.split("｜", 2)
    if len(parts) != 3 or parts[2] in {"（无样例）", ""}:
        return descriptor
    samples = " / ".join(_redact_sample(item) for item in parts[2].split(" / "))
    return "｜".join((parts[0], parts[1], samples))


def _serialize_cache(info: WorkbookInfo, key: dict[str, Any]) -> dict[str, Any]:
    preview = []
    for index, row in enumerate(info.preview_rows):
        preview.append(list(row) if index == 0 else [_redact_sample(value) for value in row])
    return {
        "key": key,
        "result": {
            "path": info.path,
            "sheets": list(info.sheets),
            "selected_sheet": info.selected_sheet,
            "header_row": info.header_row,
            "header_score": info.header_score,
            "max_row": info.max_row,
            "max_column": info.max_column,
            "actual_last_column": info.actual_last_column,
            "descriptors": [_redact_descriptor(item) for item in info.descriptors],
            "recommended_mapping": asdict(info.recommended_mapping),
            "result_columns": asdict(info.result_columns),
            # 城市全集不落检测缓存；只保存少量脱敏预览和结构化检测结果。
            "preview_rows": preview,
        },
    }


def _deserialize_cache(payload: dict[str, Any]) -> WorkbookInfo:
    result = payload["result"]
    return WorkbookInfo(
        path=result["path"],
        sheets=tuple(result["sheets"]),
        selected_sheet=result["selected_sheet"],
        header_row=int(result["header_row"]),
        header_score=int(result["header_score"]),
        max_row=int(result["max_row"]),
        max_column=int(result["max_column"]),
        actual_last_column=int(result["actual_last_column"]),
        descriptors=tuple(result["descriptors"]),
        recommended_mapping=FieldMapping(**result["recommended_mapping"]),
        result_columns=ResultColumns(**result["result_columns"]),
        city_catalog=DEFAULT_CITY_CATALOG,
        preview_rows=tuple(tuple(row) for row in result["preview_rows"]),
    )


def _load_cache(cache_dir: Path, key: dict[str, Any]) -> WorkbookInfo | None:
    target = _cache_path(cache_dir, key)
    if not target.is_file():
        return None
    try:
        payload = json.loads(target.read_text(encoding="utf-8"))
        if payload.get("key") != key:
            return None
        return _deserialize_cache(payload)
    except (OSError, ValueError, KeyError, TypeError):
        return None


def _write_cache(cache_dir: Path, key: dict[str, Any], info: WorkbookInfo) -> None:
    cache_dir.mkdir(parents=True, exist_ok=True)
    target = _cache_path(cache_dir, key)
    payload = _serialize_cache(info, key)
    handle, temporary = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp", dir=cache_dir)
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, separators=(",", ":"), default=str)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def inspect_workbook_detailed(
    path: str | Path,
    sheet_name: str | None = None,
    header_row: int | None = None,
    *,
    cache_dir: str | Path | None = None,
    cancel_event: Event | None = None,
    progress: ProgressCallback | None = None,
    timing: TimingCallback | None = None,
) -> InspectionOutcome:
    source = Path(path).resolve()
    book = None

    def notify(text: str) -> None:
        _check_cancel(cancel_event)
        if progress:
            progress(text)

    try:
        notify("正在读取工作表")

        def open_book():
            reject_xlm_macro_sheets(source)
            return load_workbook(
                source,
                read_only=True,
                data_only=False,
                keep_vba=False,
                keep_links=False,
                rich_text=False,
            )

        book = _timed("读取工作表列表", open_book, timing, details={"file": source.name})
        sheets = tuple(book.sheetnames)
        if not sheets:
            raise ValueError("工作簿不包含可读取的工作表。")
        selected = sheet_name if sheet_name in sheets else sheets[0]
        ws = book[selected]
        early_dimensions: tuple[int, int] | None = None
        if ws.max_column is None or ws.max_row is None:
            early_dimensions = _timed(
                "定位最后使用列",
                lambda: _actual_dimensions_from_ooxml(book, ws, cancel_event),
                timing,
                details={"sheet": selected, "reader": "OOXML cell references"},
            )
        effective_max_column = ws.max_column or (early_dimensions or (0, 0))[0]
        effective_max_row = ws.max_row or (early_dimensions or (0, 0))[1]
        key = _file_key(source, selected, header_row)
        if cache_dir is not None:
            cached = _timed(
                "检测缓存查找",
                lambda: _load_cache(Path(cache_dir), key),
                timing,
                details={"sheet": selected},
            )
            if cached is not None:
                return InspectionOutcome(cached, True)

        notify("正在识别表头")

        desired_sample_end = max(
            HEADER_SCAN_LIMIT,
            (header_row or HEADER_SCAN_LIMIT) + SAMPLE_DATA_ROW_LIMIT,
        )
        sample_end = (
            min(effective_max_row, desired_sample_end)
            if effective_max_row
            else desired_sample_end
        )

        def read_sample_and_header():
            rows = []
            for index, row in enumerate(
                ws.iter_rows(
                    min_row=1,
                    max_row=sample_end,
                    max_col=effective_max_column,
                    values_only=True,
                ),
                start=1,
            ):
                if index % 16 == 0:
                    _check_cancel(cancel_event)
                rows.append(tuple(row))
            detected = (
                _detect_header_from_rows(rows, HEADER_SCAN_LIMIT)
                if header_row is None
                else (header_row, _score_values(rows[header_row - 1] if header_row <= len(rows) else ()))
            )
            return rows, detected

        rows, (detected_row, detected_score) = _timed(
            "自动识别表头",
            read_sample_and_header,
            timing,
            details={"sheet": selected, "scan_limit": HEADER_SCAN_LIMIT},
        )

        notify("正在分析字段")

        def analyze_fields():
            mapping = _recommend_mapping_from_rows(rows, detected_row, effective_max_column)
            descriptors = tuple(
                _descriptor_from_rows(rows, detected_row, col)
                for col in range(1, effective_max_column + 1)
            )
            return mapping, descriptors

        mapping, descriptors = _timed(
            "生成字段候选",
            analyze_fields,
            timing,
            details={"sheet": selected, "sample_rows": SAMPLE_DATA_ROW_LIMIT},
        )
        city_catalog = _timed(
            "读取必要城市列",
            lambda: _collect_city_catalog(
                ws, mapping, detected_row, cancel_event, effective_max_row
            ),
            timing,
            details={
                "sheet": selected,
                "columns": [mapping.origin_city, mapping.destination_city],
            },
        )
        if early_dimensions is None:
            last_column, actual_max_row = _timed(
                "定位最后使用列",
                lambda: _actual_dimensions_from_ooxml(book, ws, cancel_event),
                timing,
                details={"sheet": selected, "reader": "OOXML cell references"},
            )
        else:
            last_column, actual_max_row = early_dimensions
        header_values = rows[detected_row - 1] if detected_row <= len(rows) else ()
        result_columns = _plan_result_columns_from_header(
            header_values, effective_max_column, last_column
        )

        notify("正在生成预览")

        def make_preview():
            end = min(len(rows), detected_row + PREVIEW_DATA_ROW_LIMIT)
            return tuple(
                tuple(row[: min(effective_max_column, 19)])
                for row in rows[detected_row - 1 : end]
            )

        preview = _timed(
            "生成前10行预览",
            make_preview,
            timing,
            details={"sheet": selected, "data_row_limit": PREVIEW_DATA_ROW_LIMIT},
        )
        info = WorkbookInfo(
            str(source),
            sheets,
            selected,
            detected_row,
            detected_score,
            effective_max_row or actual_max_row,
            effective_max_column,
            last_column,
            descriptors,
            mapping,
            result_columns,
            city_catalog,
            preview,
        )
        if cache_dir is not None:
            _timed(
                "写入检测缓存",
                lambda: _write_cache(Path(cache_dir), key, info),
                timing,
                details={"sheet": selected, "redacted": True},
            )
        return InspectionOutcome(info, False)
    finally:
        if book is not None:
            book.close()


def inspect_workbook(
    path: str | Path, sheet_name: str | None = None, header_row: int | None = None
) -> WorkbookInfo:
    return inspect_workbook_detailed(path, sheet_name, header_row).info


def iter_row_inputs(selection: Any):
    book = load_workbook(
        selection.path,
        read_only=True,
        data_only=False,
        keep_vba=False,
        keep_links=False,
        rich_text=False,
    )
    try:
        ws = book[selection.sheet_name]
        fields = selection.fields
        selected_columns = [
            fields.quote_type,
            fields.origin_city,
            fields.origin_address,
            fields.destination_city,
            fields.destination_address,
        ]
        if fields.vehicle is not None:
            selected_columns.append(fields.vehicle)
        first_col = min(selected_columns)
        last_col = max(selected_columns)

        def value(values: tuple[Any, ...], column: int | None) -> Any:
            return None if column is None else values[column - first_col]

        rows = ws.iter_rows(
            min_row=selection.header_row + 1,
            max_row=ws.max_row,
            min_col=first_col,
            max_col=last_col,
            values_only=True,
        )
        for excel_row, values in enumerate(rows, start=selection.header_row + 1):
            yield RowInput(
                excel_row,
                value(values, fields.quote_type),
                value(values, fields.origin_city),
                value(values, fields.origin_address),
                value(values, fields.destination_city),
                value(values, fields.destination_address),
                value(values, fields.vehicle),
            )
    finally:
        book.close()
