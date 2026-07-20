from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.addressing.service import AddressBookService, TRUSTED_STATUSES
from src.domain.models import AddressConfirmationStatus, ConfirmedAddress
from src.utils.text import normalized_text


PENDING_HEADERS = (
    "地址ID",
    "原始城市",
    "原始详细地址",
    "高德标准地址",
    "定位层级",
    "修正查询地址",
    "是否确认",
    "确认备注",
)


@dataclass(frozen=True)
class ImportSummary:
    success: int
    skipped: int
    failed: int
    errors: tuple[str, ...] = ()


def export_pending_addresses(path: str | Path, addresses: list[ConfirmedAddress]) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待确认地址"
    sheet.append(PENDING_HEADERS)
    pending = [item for item in addresses if item.confirmation_status not in TRUSTED_STATUSES]
    for item in pending:
        sheet.append((
            item.address_id,
            item.original_city,
            item.original_address,
            item.formatted_address,
            item.level,
            item.query_address,
            "",
            item.confirmation_note,
        ))
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = (22, 12, 42, 42, 12, 42, 12, 30)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{max(1, sheet.max_row)}"
    if sheet.max_row >= 2:
        validation = DataValidation(type="list", formula1='"是,否,地址无效,待进一步核实"')
        sheet.add_data_validation(validation)
        validation.add(f"G2:G{sheet.max_row}")
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(target)
    workbook.close()
    return target


def import_confirmation_results(
    path: str | Path,
    service: AddressBookService,
    *,
    confirmed_by: str = "Excel 导入",
) -> ImportSummary:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    errors: list[str] = []
    success = skipped = failed = 0
    try:
        sheet = workbook[workbook.sheetnames[0]]
        headers = [normalized_text(sheet.cell(1, column).value) for column in range(1, sheet.max_column + 1)]
        positions = {header: index + 1 for index, header in enumerate(headers)}
        missing = [header for header in PENDING_HEADERS if header not in positions]
        if missing:
            return ImportSummary(0, 0, 1, (f"缺少列：{'、'.join(missing)}",))

        seen: dict[str, tuple[str, str, str]] = {}
        for row in range(2, sheet.max_row + 1):
            address_id = normalized_text(sheet.cell(row, positions["地址ID"]).value)
            if not address_id:
                skipped += 1
                continue
            query = normalized_text(sheet.cell(row, positions["修正查询地址"]).value)
            decision = normalized_text(sheet.cell(row, positions["是否确认"]).value, remove_all_space=True)
            note = normalized_text(sheet.cell(row, positions["确认备注"]).value)
            signature = (query, decision, note)
            if address_id in seen:
                if seen[address_id] != signature:
                    failed += 1
                    errors.append(f"第 {row} 行：地址ID重复且内容冲突")
                else:
                    skipped += 1
                continue
            seen[address_id] = signature
            current = service.repository.get_address(address_id)
            if current is None:
                failed += 1
                errors.append(f"第 {row} 行：地址ID不存在")
                continue
            try:
                if decision in {"", "否"}:
                    if note != current.confirmation_note:
                        service.repository.put_address(replace(current, confirmation_note=note))
                        success += 1
                    else:
                        skipped += 1
                    continue
                if decision == "地址无效":
                    service.mark_invalid(address_id, note=note)
                    success += 1
                    continue
                if decision == "待进一步核实":
                    service.repository.put_address(replace(
                        current,
                        confirmation_status=AddressConfirmationStatus.NEEDS_FURTHER_REVIEW.value,
                        confirmation_note=note,
                    ))
                    success += 1
                    continue
                if decision != "是":
                    failed += 1
                    errors.append(f"第 {row} 行：是否确认值无效“{decision}”")
                    continue
                if query and query != current.query_address:
                    service.save_query_address(address_id, query)
                    if service.geocoder is None:
                        failed += 1
                        errors.append(f"第 {row} 行：修正地址已保存，但缺少地理编码客户端，尚未确认")
                        continue
                    service.reparse(address_id, confirm_corrected=True)
                    reparsed = service.repository.get_address(address_id)
                    if reparsed and note:
                        service.repository.put_address(replace(
                            reparsed, confirmed_by=confirmed_by, confirmation_note=note
                        ))
                else:
                    service.confirm(address_id, confirmed_by=confirmed_by, note=note)
                success += 1
            except Exception as exc:
                failed += 1
                errors.append(f"第 {row} 行：{exc}")
    finally:
        workbook.close()
    return ImportSummary(success, skipped, failed, tuple(errors))
