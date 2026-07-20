from __future__ import annotations

import json
import shutil
import traceback
import uuid
from datetime import datetime
from pathlib import Path
from typing import Callable

from openpyxl.utils import get_column_letter

from src.amap.http_client import AmapHttpClient, RealApiAuditLogger
from src.amap.driving_clients import DRIVING_MODE, RealDrivingRouter
from src.amap.real_clients import RealGeocoder
from src.addressing.service import AddressBookService
from src.application.processor import BusinessProcessor
from src.application.task_control import TaskController
from src.cache.sqlite_cache import CacheRepository
from src.domain.models import (
    OfficeEngine,
    ProgressSnapshot,
    RunSummary,
    TaskMode,
    TaskState,
    WorkbookSelection,
)
from src.office.backend import OfficeBackend
from src.security.key_store import SecureKeyStore
from src.utils.files import assert_unchanged, atomic_promote, fingerprint, is_within, reject_xlm_macro_sheets
from src.utils.text import stable_hash
from src.validation.quote import is_target_quote
from src.validation.address import CLEANER_VERSION
from src.workbook.preview import inspect_workbook, iter_row_inputs


class PrototypeRunner:
    def __init__(self, project_root: str | Path) -> None:
        self.root = Path(project_root).resolve()
        self.cache_path = self.root / "cache" / "task_003b.sqlite"
        self.runtime_root = self.root / "samples" / "working" / "prototype"
        self.log_root = self.root / "logs" / "prototype"
        self.controller = TaskController()
        self._repository: CacheRepository | None = None
        self._active_task_id: str | None = None

    def pause(self) -> bool:
        accepted = self.controller.request_pause()
        if accepted and self._repository and self._active_task_id:
            self._repository.update_task(self._active_task_id, TaskState.PAUSING.value)
        return accepted

    def resume(self) -> bool:
        accepted = self.controller.resume()
        if accepted and self._repository and self._active_task_id:
            self._repository.update_task(self._active_task_id, TaskState.RUNNING.value)
        return accepted

    def stop(self) -> bool:
        accepted = self.controller.request_stop()
        if accepted and self._repository and self._active_task_id:
            self._repository.update_task(self._active_task_id, TaskState.STOPPING.value)
        return accepted

    @staticmethod
    def default_output_name(source: str | Path, mode: TaskMode = TaskMode.MOCK) -> str:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        if mode == TaskMode.DRIVING_REAL:
            return f"{Path(source).stem}_普通驾车距离结果_待验收_{timestamp}.xlsm"
        return f"{Path(source).stem}_货车距离模拟结果_{timestamp}.xlsm"

    def run(
        self,
        selection: WorkbookSelection,
        *,
        engine: OfficeEngine,
        mode: TaskMode,
        output_dir: str | Path,
        final_path: str | Path | None = None,
        task_id: str | None = None,
        force_all: bool = False,
        force_route_refresh: bool = False,
        progress: Callable[[ProgressSnapshot], None] | None = None,
        checkpoint_batch: int = 10,
        checkpoint_seconds: float = 60.0,
    ) -> RunSummary:
        source = Path(selection.path).resolve()
        reject_xlm_macro_sheets(source)
        input_root = self.root / "samples" / "input"
        output_directory = Path(output_dir).resolve()
        if is_within(output_directory, input_root):
            raise ValueError("输出目录不能位于 samples/input；原始样表目录只读。")
        output_directory.mkdir(parents=True, exist_ok=True)
        final = Path(final_path).resolve() if final_path else output_directory / self.default_output_name(source, mode)
        if final == source or is_within(final, input_root):
            raise ValueError("输出文件不能覆盖输入文件或写入 samples/input。")
        info = inspect_workbook(source, selection.sheet_name, selection.header_row)
        before = fingerprint(source)
        if task_id is None:
            if force_all:
                task_id = datetime.now().strftime("%Y%m%d%H%M%S") + "_" + uuid.uuid4().hex[:8]
            else:
                task_id = "incremental_" + stable_hash(
                    before.sha256,
                    str(source),
                    selection.sheet_name,
                    selection.header_row,
                    selection.fields.as_dict(),
                    mode.value,
                    engine.value,
                )[:20]
        if mode == TaskMode.DRIVING_REAL:
            key_store = SecureKeyStore(self.root)
            if not key_store.get_key():
                raise RuntimeError("API Key 未配置：请在本机原型的“API 设置”中输入，不要在聊天中发送 Key。")
            cache_path = self.root / "cache" / "driving_real.sqlite"
            runtime_root = self.root / "samples" / "working" / "driving_real"
            log_root = self.root / "logs" / "driving_real"
        else:
            key_store = None
            cache_path = self.cache_path
            runtime_root = self.runtime_root
            log_root = self.log_root
        task_runtime = runtime_root / task_id
        task_runtime.mkdir(parents=True, exist_ok=True)
        working_copy = task_runtime / source.name
        if not working_copy.exists():
            shutil.copy2(source, working_copy)
        if fingerprint(working_copy).sha256 != before.sha256:
            raise RuntimeError("工作副本与输入文件哈希不一致，已停止。")
        partial = output_directory / f".{final.stem}.{task_id}.partial.xlsm"
        log_root.mkdir(parents=True, exist_ok=True)
        repository = CacheRepository(cache_path)
        self._repository = repository
        self._active_task_id = task_id
        office = OfficeBackend(engine, task_runtime, log_root / task_id)
        repository.create_task(
            task_id,
            before.sha256,
            {
                "path": str(source), "sheet_name": selection.sheet_name,
                "header_row": selection.header_row, "fields": selection.fields.as_dict(),
            },
            mode.value,
            str(partial),
        )
        existing_task = repository.get_task(task_id)
        if existing_task and existing_task["source_sha256"] != before.sha256:
            raise RuntimeError("恢复任务的源文件哈希与当前输入不一致，已停止。")
        resume_after = int(existing_task["save_watermark"]) if existing_task and partial.exists() and not force_all else 0
        repository.update_task(task_id, TaskState.RUNNING.value, output_path=str(partial))
        first_save = not partial.exists()
        last_saved_count = 0
        outcomes = []
        run_log: dict[str, object] = {
            "task_id": task_id, "source_before": before.as_dict(), "working_copy": str(working_copy),
            "partial_path": str(partial), "final_path": str(final), "engine": engine.value,
            "mode": mode.value, "checkpoint_batch": checkpoint_batch, "checkpoint_seconds": checkpoint_seconds,
        }
        http: AmapHttpClient | None = None

        def save_checkpoint(current, forced: bool) -> None:
            nonlocal first_save, last_saved_count
            if not current and not partial.exists():
                return
            source_for_save = working_copy if first_save else partial
            office.write_checkpoint(
                source_path=source_for_save, output_path=partial, first_save=first_save,
                sheet_name=selection.sheet_name, header_row=selection.header_row,
                field_mapping=selection.fields, result_columns=info.result_columns,
                outcomes=current, last_data_row=info.max_row,
            )
            first_save = False
            last_saved_count = len(current)
            state = self.controller.state
            if state == TaskState.PAUSING:
                db_state = TaskState.PAUSED.value
            elif state == TaskState.STOPPING:
                db_state = TaskState.STOPPED.value
            else:
                db_state = TaskState.RUNNING.value
            repository.update_task(task_id, db_state, output_path=str(partial), save_watermark=last_saved_count)
            if progress:
                progress(
                    ProgressSnapshot("Office 安全保存" if not forced else "Office 强制保存", 0, 0,
                                     len(current), 0, sum(bool(item.warning_fields) for item in current), 0,
                                     sum(item.cache_reused for item in current), current[-1].excel_row if current else None)
                )

        try:
            rows = list(iter_row_inputs(selection))
            target_total = sum(is_target_quote(row.quote_type) for row in rows)
            if mode == TaskMode.DRIVING_REAL:
                http = AmapHttpClient(
                    key_store,
                    audit_logger=RealApiAuditLogger(log_root / task_id / "http"),
                    max_attempts=3,
                )
                processor = BusinessProcessor(
                    repository,
                    geocoder=RealGeocoder(
                        http, cleaner_version=CLEANER_VERSION, mode=DRIVING_MODE
                    ),
                    router=RealDrivingRouter(http, cleaner_version=CLEANER_VERSION),
                )
                AddressBookService(repository, geocoder=processor.geocoder).sync_rows(
                    rows, info.city_catalog
                )
            else:
                processor = BusinessProcessor(repository)
            outcomes = processor.process(
                rows, task_id=task_id, mode=mode, city_catalog=info.city_catalog,
                force_all=force_all, force_route_refresh=force_route_refresh,
                controller=self.controller, progress=progress,
                checkpoint=save_checkpoint, checkpoint_batch=checkpoint_batch,
                checkpoint_seconds=checkpoint_seconds, resume_after=resume_after,
            )
            if self.controller.state == TaskState.STOPPED:
                repository.update_task(task_id, TaskState.STOPPED.value, output_path=str(partial), save_watermark=last_saved_count)
                assert_unchanged(before, source)
                summary = RunSummary(
                    task_id, TaskState.STOPPED.value, str(source), str(partial) if partial.exists() else None,
                    target_total, len(outcomes), sum(item.distance_km is not None for item in outcomes),
                    sum(bool(item.warning_fields) for item in outcomes), sum(item.cache_reused for item in outcomes), True,
                    actual_http_calls=http.actual_calls if http else 0,
                    unique_route_count=len({item.audit.get("route_cache_key") for item in outcomes if item.audit.get("route_cache_key")}),
                    multi_destination_count=sum(item.status == "多目的地待确认" for item in outcomes),
                    city_conflict_count=sum("地址冲突待确认" in item.status for item in outcomes),
                )
            else:
                if not partial.exists():
                    save_checkpoint(outcomes, True)
                field_columns = {**selection.fields.as_dict(), **info.result_columns.as_dict()}
                warning_cells = sorted({
                    f"{get_column_letter(field_columns[field])}{item.excel_row}"
                    for item in outcomes for field in item.warning_fields
                    if field in field_columns and field_columns[field] is not None
                })
                verification = office.verify(
                    path=partial, sheet_name=selection.sheet_name, header_row=selection.header_row,
                    result_columns=info.result_columns, target_rows=[item.excel_row for item in outcomes],
                    last_data_row=info.max_row, warning_cells=warning_cells,
                )
                verify_result = verification["result"]
                if verify_result["file_format"] != 52 or verify_result["status_count"] != target_total:
                    raise RuntimeError(f"Office 重开验证未通过：{verify_result}")
                assert_unchanged(before, source)
                atomic_promote(partial, final)
                self.controller.mark_completed()
                repository.update_task(task_id, TaskState.COMPLETED.value, output_path=str(final), save_watermark=len(outcomes))
                summary = RunSummary(
                    task_id, TaskState.COMPLETED.value, str(source), str(final), target_total, len(outcomes),
                    sum(item.distance_km is not None for item in outcomes), sum(bool(item.warning_fields) for item in outcomes),
                    sum(item.cache_reused for item in outcomes), False,
                    actual_http_calls=http.actual_calls if http else 0,
                    unique_route_count=len({item.audit.get("route_cache_key") for item in outcomes if item.audit.get("route_cache_key")}),
                    multi_destination_count=sum(item.status == "多目的地待确认" for item in outcomes),
                    city_conflict_count=sum("地址冲突待确认" in item.status for item in outcomes),
                )
                run_log["office_reopen_verification"] = verification
            run_log["summary"] = summary.__dict__
            run_log["source_after"] = fingerprint(source).as_dict()
            return summary
        except Exception as exc:
            self.controller.mark_failed()
            repository.update_task(task_id, TaskState.FAILED.value, output_path=str(partial), save_watermark=last_saved_count, error=repr(exc))
            run_log.update({"error": repr(exc), "traceback": traceback.format_exc(), "source_after": fingerprint(source).as_dict()})
            raise
        finally:
            log_path = log_root / task_id / "run_summary.json"
            log_path.parent.mkdir(parents=True, exist_ok=True)
            log_path.write_text(json.dumps(run_log, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
            repository.close()
            self._repository = None
            self._active_task_id = None
