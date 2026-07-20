from __future__ import annotations

from copy import copy

from openpyxl import Workbook

from src.domain.models import ResultColumns
from src.workbook.preview import inspect_workbook, plan_result_columns


def test_sample_dynamic_recommendation(sample_path):
    info = inspect_workbook(sample_path)
    assert info.header_row == 1
    assert info.recommended_mapping.as_dict() == {
        "quote_type": 13, "origin_city": 5, "origin_address": 6,
        "destination_city": 7, "destination_address": 8, "vehicle": 11,
    }
    assert info.result_columns == ResultColumns(17, 18, 19)


def test_descriptors_have_letter_header_sample(sample_path):
    info = inspect_workbook(sample_path)
    assert info.descriptors[12].startswith("M｜报价类型｜")
    assert "非线路报价" in info.descriptors[12]


def test_header_detection_with_shifted_layout(tmp_path):
    path = tmp_path / "shifted.xlsx"
    book = Workbook()
    ws = book.active
    ws.append(["说明"])
    ws.append(["始发仓", "", "到达仓", "", "车型", "报价类型"])
    ws.append(["广州", "广东省广州市A路1号", "深圳", "广东省深圳市B路2号", 4.2, "非线路报价"])
    book.save(path)
    info = inspect_workbook(path)
    assert info.header_row == 2
    assert info.recommended_mapping.quote_type == 6
    assert info.recommended_mapping.origin_address == 2


def test_result_columns_reuse_existing_headers():
    book = Workbook()
    ws = book.active
    ws.append(["报价类型", "高德普通驾车公路距离（公里）", "距离查询状态", "距离查询说明"])
    assert plan_result_columns(ws, 1) == ResultColumns(2, 3, 4)


def test_old_truck_result_columns_are_preserved_and_new_triplet_is_appended():
    book = Workbook()
    ws = book.active
    ws.append(["报价类型", "高德货车公路距离（公里）", "距离查询状态", "距离查询说明"])
    assert plan_result_columns(ws, 1) == ResultColumns(5, 6, 7)


def test_non_contiguous_driving_headers_are_not_reused():
    book = Workbook()
    ws = book.active
    ws.append(["高德普通驾车公路距离（公里）", "旧列", "距离查询状态", "距离查询说明"])
    assert plan_result_columns(ws, 1) == ResultColumns(5, 6, 7)


def test_vehicle_mapping_is_optional(tmp_path):
    path = tmp_path / "no_vehicle.xlsx"
    book = Workbook()
    ws = book.active
    ws.append(["始发城市", "始发详细地址", "目的城市", "目的详细地址", "报价类型"])
    ws.append(["广州", "A路", "深圳", "B路", "非线路报价"])
    book.save(path)
    info = inspect_workbook(path)
    assert info.recommended_mapping.vehicle is None


def test_result_columns_ignore_style_only_cells():
    book = Workbook()
    ws = book.active
    ws["A1"] = "报价类型"
    ws["J10"].fill = copy(ws["A1"].fill)
    assert plan_result_columns(ws, 1) == ResultColumns(2, 3, 4)
