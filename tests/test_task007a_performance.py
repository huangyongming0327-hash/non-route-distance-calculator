from __future__ import annotations

import json
import os
import time
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog

from src.performance import PerformanceLogger
from src.ui.main_window import MainWindow, OutputPermissionWorker
from src.workbook.preview import inspect_workbook_detailed


def _make_workbook(path: Path, data_rows: int, header_row: int) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    sheet = book.create_sheet("运距 数据")
    for row_number in range(1, header_row):
        sheet.append([f"说明{row_number}"])
    sheet.append([
        "始发城市", "始发详细地址", "目的城市", "目的详细地址", "车型", "报价类型"
    ])
    for index in range(data_rows):
        sheet.append([
            "广州", f"广东省广州市测试路{index}号", "深圳",
            f"广东省深圳市样例街{index}号", "4.2米", "非线路报价",
        ])
    book.save(path)
    return path


@pytest.mark.parametrize(
    ("data_rows", "header_row", "directory", "filename"),
    [
        (50, 1, "普通路径", "header1.xlsx"),
        (2_000, 4, "中文路径", "两千行 样表.xlsx"),
        (10_000, 15, "带 空格 路径", "一万行 样表.xlsx"),
    ],
)
def test_fast_readonly_detection_sizes_headers_and_paths(
    tmp_path: Path,
    data_rows: int,
    header_row: int,
    directory: str,
    filename: str,
) -> None:
    path = _make_workbook(tmp_path / directory / filename, data_rows, header_row)
    started = time.perf_counter()
    outcome = inspect_workbook_detailed(path, cache_dir=tmp_path / "检测缓存")
    elapsed = time.perf_counter() - started
    assert outcome.info.header_row == header_row
    assert outcome.info.max_row == data_rows + header_row
    assert outcome.info.recommended_mapping.quote_type == 6
    assert elapsed < 5


def test_current_281_row_sample_and_cache_hit_under_one_second(
    sample_path: Path, tmp_path: Path
) -> None:
    cache_dir = tmp_path / "cache"
    first = inspect_workbook_detailed(sample_path, cache_dir=cache_dir)
    started = time.perf_counter()
    second = inspect_workbook_detailed(sample_path, cache_dir=cache_dir)
    elapsed = time.perf_counter() - started
    assert first.info.max_row == 281
    assert not first.cache_hit
    assert second.cache_hit
    assert elapsed < 1


def test_detection_cache_key_invalidation_and_redaction(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "缓存测试.xlsx", 5, 1)
    cache_dir = tmp_path / "cache"
    first = inspect_workbook_detailed(path, cache_dir=cache_dir)
    second = inspect_workbook_detailed(path, cache_dir=cache_dir)
    assert not first.cache_hit and second.cache_hit

    payload = json.loads(next(cache_dir.glob("*.json")).read_text(encoding="utf-8"))
    assert set(payload["key"]) >= {
        "canonical_path", "file_size", "modified_time_ns", "sheet_name", "algorithm_version"
    }
    cached_text = json.dumps(payload, ensure_ascii=False)
    assert "广东省广州市测试路0号" not in cached_text
    assert "广东省深圳市样例街0号" not in cached_text

    book = load_workbook(path)
    book.active["G2"] = "文件已变化"
    book.save(path)
    stat = path.stat()
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
    changed = inspect_workbook_detailed(path, cache_dir=cache_dir)
    assert not changed.cache_hit


def test_header_detection_does_not_launch_office_or_http(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "offline.xlsx", 20, 1)
    with (
        patch("subprocess.Popen", side_effect=AssertionError("不得启动 Office 进程")),
        patch(
            "src.amap.http_client.StandardLibraryTransport.get",
            side_effect=AssertionError("不得调用真实高德 API"),
        ),
    ):
        result = inspect_workbook_detailed(path)
    assert result.info.header_row == 1


@pytest.mark.parametrize("file_count", [0, 1_000])
def test_output_permission_check_empty_or_large_directory(
    tmp_path: Path, file_count: int
) -> None:
    directory = tmp_path / f"output_{file_count}"
    directory.mkdir()
    for index in range(file_count):
        (directory / f"existing_{index}.txt").touch()
    logger = PerformanceLogger(tmp_path / "performance.jsonl")
    worker = OutputPermissionWorker(1, str(directory), logger, timeout_seconds=0.5)
    result: list[tuple[int, str, str]] = []
    worker.finished.connect(lambda *values: result.append(values))
    started = time.perf_counter()
    worker.run()
    assert time.perf_counter() - started < 0.5
    assert result[0][1] == "writable"
    assert len(list(directory.iterdir())) == file_count


def test_output_permission_slow_disk_times_out_without_blocking_wait(tmp_path: Path) -> None:
    logger = PerformanceLogger(tmp_path / "performance.jsonl")
    worker = OutputPermissionWorker(1, str(tmp_path), logger, timeout_seconds=0.03)
    result: list[tuple[int, str, str]] = []
    worker.finished.connect(lambda *values: result.append(values))

    def slow_probe() -> tuple[str, str]:
        time.sleep(0.2)
        return "writable", "可写"

    worker._probe = slow_probe  # type: ignore[method-assign]
    started = time.perf_counter()
    worker.run()
    assert time.perf_counter() - started < 0.15
    assert result[0][1] == "timeout"


def _wait_until(app: QApplication, predicate, timeout_ms: int = 5_000) -> None:
    deadline = time.perf_counter() + timeout_ms / 1000
    while time.perf_counter() < deadline:
        app.processEvents()
        if predicate():
            return
        QTest.qWait(10)
    raise AssertionError("等待后台操作完成超时")


def test_detection_thread_tab_switch_duplicate_click_and_file_change(
    project_root: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    first = _make_workbook(tmp_path / "first.xlsx", 2_000, 4)
    second = _make_workbook(tmp_path / "second.xlsx", 50, 15)
    from src.ui import main_window as main_window_module

    real_inspect = main_window_module.inspect_workbook_detailed

    def delayed_inspect(*args, cancel_event=None, progress=None, **kwargs):
        if progress:
            progress("正在读取工作表")
        for _ in range(20):
            if cancel_event is not None and cancel_event.is_set():
                from src.workbook.preview import InspectionCancelled

                raise InspectionCancelled("用户已取消当前检测。")
            time.sleep(0.005)
        return real_inspect(
            *args, cancel_event=cancel_event, progress=progress, **kwargs
        )

    monkeypatch.setattr(main_window_module, "inspect_workbook_detailed", delayed_inspect)
    app = QApplication.instance() or QApplication([])
    window = MainWindow(project_root, defer_initial_load=True)
    window.show()
    try:
        window.file_edit.setText(str(first))
        window._inspect_current(force_auto_header=True)
        active_id = window.inspection_request_sequence
        assert window.cancel_detect_button.isEnabled()
        assert window.detect_button.text() == "检测中…"

        # 快速重复请求被拒绝；同时页签仍可切换和处理事件。
        window._inspect_current(force_auto_header=True)
        assert window.inspection_request_sequence == active_id
        window.tab_widget.setCurrentIndex(2)
        app.processEvents()
        assert window.tab_widget.currentIndex() == 2

        # 更换文件会取消旧任务并排队新任务，旧结果不得覆盖新文件。
        window.file_edit.setText(str(second))
        window._inspect_current(force_auto_header=True)
        _wait_until(
            app,
            lambda: window.info is not None
            and Path(window.info.path) == second.resolve()
            and window.startup_thread is None,
        )
        assert window.info.header_row == 15
        assert window.detect_button.text() == "自动检测表头/字段"
    finally:
        window.close()
        _wait_until(app, lambda: not window.isVisible())


def test_select_output_returns_before_background_permission_check(
    project_root: Path, tmp_path: Path
) -> None:
    app = QApplication.instance() or QApplication([])
    window = MainWindow(project_root, defer_initial_load=True)
    selected = tmp_path / "含大量文件的输出目录"
    selected.mkdir()
    calls: list[str] = []
    try:
        with (
            patch.object(QFileDialog, "getExistingDirectory", return_value=str(selected)),
            patch.object(
                window,
                "_start_output_permission_check",
                side_effect=lambda path: calls.append(path),
            ),
            patch.object(
                window,
                "_inspect_current",
                side_effect=AssertionError("选择输出目录不得分析 Excel"),
            ),
            patch.object(
                window,
                "_refresh_office_engines",
                side_effect=AssertionError("选择输出目录不得执行 Office 检测"),
            ),
        ):
            window._select_output()
            assert window.output_edit.text() == str(selected)
            app.processEvents()
        assert calls == [str(selected)]
    finally:
        window.close()
        app.processEvents()
