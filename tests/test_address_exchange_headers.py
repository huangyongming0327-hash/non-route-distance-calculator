from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.addressing.exchange import (
    PENDING_HEADERS,
    export_pending_addresses,
    import_confirmation_results,
)
from src.addressing.service import AddressBookService
from src.cache.sqlite_cache import CacheRepository
from src.domain.models import AddressConfirmationStatus, RowInput


@pytest.fixture
def pending_address(tmp_path):
    repository = CacheRepository(tmp_path / "addresses.db")
    service = AddressBookService(repository)
    records, _ = service.sync_rows(
        [RowInput(2, "非线路报价", "广州", "测试仓库1号", "深圳", "目标仓库2号", None)],
        ["广州", "深圳"],
    )
    record = next(item for item in records if item.original_city == "广州")
    try:
        yield service, record
    finally:
        repository.close()


def _write_import_file(
    path: Path,
    record,
    *,
    header_row: int,
    headers: tuple[str, ...] = PENDING_HEADERS,
    leading_values: tuple[str | None, ...] = (),
    decision: str = "待进一步核实",
    note: str = "表头自动识别测试",
) -> Path:
    values = {
        "地址ID": record.address_id,
        "原始城市": record.original_city,
        "原始详细地址": record.original_address,
        "高德标准地址": record.formatted_address,
        "定位层级": record.level,
        "修正查询地址": record.query_address,
        "是否确认": decision,
        "确认备注": note,
    }
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待确认地址"
    for row_number in range(1, header_row):
        value = leading_values[row_number - 1] if row_number <= len(leading_values) else None
        sheet.append([value])
    sheet.append(headers)
    sheet.append([values.get(header, "") for header in headers])
    workbook.save(path)
    workbook.close()
    return path


def test_import_detects_header_in_first_row(tmp_path, pending_address) -> None:
    service, record = pending_address
    target = _write_import_file(tmp_path / "row1.xlsx", record, header_row=1)

    summary = import_confirmation_results(target, service)

    assert (summary.success, summary.skipped, summary.failed) == (1, 0, 0)
    updated = service.repository.get_address(record.address_id)
    assert updated.address_id == record.address_id
    assert updated.original_city == record.original_city
    assert updated.original_address == record.original_address


def test_import_detects_reordered_header_in_fourth_row(tmp_path, pending_address) -> None:
    service, record = pending_address
    reordered = tuple(reversed(PENDING_HEADERS))
    target = _write_import_file(
        tmp_path / "row4_reordered.xlsx",
        record,
        header_row=4,
        headers=reordered,
        leading_values=("TASK-005A 唯一地址待确认清单", "说明文字", None),
    )

    summary = import_confirmation_results(target, service)

    assert (summary.success, summary.skipped, summary.failed) == (1, 0, 0)
    assert service.repository.get_address(record.address_id).confirmation_note == "表头自动识别测试"


def test_import_ignores_merged_title_before_header(tmp_path, pending_address) -> None:
    service, record = pending_address
    target = _write_import_file(
        tmp_path / "merged_title.xlsx",
        record,
        header_row=3,
        leading_values=("TASK-005A 唯一地址待确认清单", "按地址ID导回"),
    )
    workbook = load_workbook(target)
    sheet = workbook["待确认地址"]
    sheet.merge_cells("A1:H1")
    workbook.save(target)
    workbook.close()

    summary = import_confirmation_results(target, service)

    assert (summary.success, summary.skipped, summary.failed) == (1, 0, 0)


def test_import_ignores_instructions_and_blank_rows(tmp_path, pending_address) -> None:
    service, record = pending_address
    target = _write_import_file(
        tmp_path / "instructions_and_blanks.xlsx",
        record,
        header_row=5,
        leading_values=(
            "填写说明：原Excel地址不会被修改",
            None,
            "“是否确认”可填写是、否、地址无效或待进一步核实",
            None,
        ),
    )

    summary = import_confirmation_results(target, service)

    assert (summary.success, summary.skipped, summary.failed) == (1, 0, 0)


def test_import_reports_scan_summary_and_candidate_for_missing_field(
    tmp_path,
    pending_address,
) -> None:
    service, record = pending_address
    incomplete = tuple(header for header in PENDING_HEADERS if header != "确认备注")
    target = _write_import_file(
        tmp_path / "missing_header.xlsx",
        record,
        header_row=4,
        headers=incomplete,
        leading_values=("TASK-005A 唯一地址待确认清单", "说明", None),
    )

    summary = import_confirmation_results(target, service)

    assert (summary.success, summary.skipped, summary.failed) == (0, 0, 1)
    message = summary.errors[0]
    assert "前 20 行" in message
    assert "实际扫描" in message
    assert "TASK-005A 唯一地址待确认清单" in message
    assert "可能的表头行：第 4 行" in message
    assert "匹配 7/8" in message
    assert "缺少：确认备注" in message


def test_program_export_can_be_imported_directly_without_deleting_rows(
    tmp_path,
    pending_address,
) -> None:
    service, record = pending_address
    target = export_pending_addresses(tmp_path / "program_export.xlsx", [record])

    summary = import_confirmation_results(target, service)

    assert summary.failed == 0
    assert summary.success == 0
    assert summary.skipped == 1
    unchanged = service.repository.get_address(record.address_id)
    assert unchanged.address_id == record.address_id
    assert unchanged.original_address == record.original_address
    assert unchanged.confirmation_status == AddressConfirmationStatus.UNCONFIRMED.value
