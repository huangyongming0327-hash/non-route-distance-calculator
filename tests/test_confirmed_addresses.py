from __future__ import annotations

from dataclasses import replace

from openpyxl import load_workbook

from src.addressing.exchange import export_pending_addresses, import_confirmation_results
from src.addressing.service import AddressBookService, address_id_for
from src.cache.sqlite_cache import CacheRepository
from src.domain.models import (
    AddressConfirmationStatus,
    AddressRiskLevel,
    GeocodeResult,
    RouteResult,
    RowInput,
    TaskMode,
)
from src.application.processor import BusinessProcessor
from src.utils.text import stable_hash


class CountingGeocoder:
    mode = "driving_real"
    cache_contract = "test-geocode-v1"

    def __init__(self, *, level: str = "门牌号") -> None:
        self.calls = 0
        self.level = level

    def cache_key_for(self, address: str, city: str = "") -> str:
        return stable_hash("test-geocode", address, city)

    def geocode(self, address: str, city: str = "") -> GeocodeResult:
        self.calls += 1
        suffix = self.calls
        return GeocodeResult(
            self.cache_key_for(address, city), address, city,
            113.000000 + suffix / 1000, 23.000000 + suffix / 1000,
            self.level, f"广东省{city}测试区{address}", mode=self.mode,
            province="广东省", city=f"{city}市" if city and not city.endswith("市") else city,
            district="测试区", street=address if address.endswith("路") else "测试路",
            number="1号", api_version="test-geocode-v1", cleaner_version="address-cleaner-v1",
        )


class CountingRouter:
    mode = "driving_real"
    cache_contract = "test-route-v1"

    def __init__(self) -> None:
        self.calls = 0

    def cache_key_for(self, origin, destination) -> str:
        return stable_hash(
            "test-route", origin.geocode_version or origin.cache_key,
            destination.geocode_version or destination.cache_key,
        )

    def route(self, origin, destination) -> RouteResult:
        self.calls += 1
        return RouteResult(
            self.cache_key_for(origin, destination), 10.0, mode=self.mode,
            raw_distance=10000, path_count=1, api_contract_version="test-route-v1",
        )


def row(excel_row: int = 2, *, origin_city="广州", origin="测试路", destination_city="深圳", destination="目标路"):
    return RowInput(excel_row, "非线路报价", origin_city, origin, destination_city, destination, None)


def test_confirmed_address_crud_and_route_invalidation(tmp_path):
    geo = CountingGeocoder()
    with CacheRepository(tmp_path / "cache.db") as repository:
        service = AddressBookService(repository, geocoder=geo)
        record = service.resolve("广州", "测试路", city_conflict=False).record
        assert repository.address_count() == 1
        route = RouteResult("route", 1.0, mode="driving_real")
        repository.put_route(
            route, "o", "d", origin_address_id=record.address_id,
            origin_geocode_version=record.geocode_version,
        )
        changed = service.save_query_address(record.address_id, "测试路1号")
        assert changed.confirmation_status == AddressConfirmationStatus.UNCONFIRMED.value
        assert repository.get_route("route") is None


def test_same_address_reused_across_rows_and_workbooks(tmp_path):
    geo = CountingGeocoder()
    path = tmp_path / "cache.db"
    with CacheRepository(path) as repository:
        first = AddressBookService(repository, geocoder=geo).resolve("广州", "测试路", city_conflict=False)
        second = AddressBookService(repository, geocoder=geo).resolve("广州", "测试路", city_conflict=False)
        assert first.record.address_id == second.record.address_id
        assert second.confirmed_reused is True
        assert geo.calls == 1
    with CacheRepository(path) as repository:
        third = AddressBookService(repository, geocoder=geo).resolve("广州", "测试路", city_conflict=False)
        assert third.confirmed_reused is True
        assert geo.calls == 1


def test_corrected_address_regeocodes_and_confirms(tmp_path):
    geo = CountingGeocoder()
    with CacheRepository(tmp_path / "cache.db") as repository:
        service = AddressBookService(repository, geocoder=geo)
        first = service.resolve("广州", "测试路", city_conflict=False).record
        service.save_query_address(first.address_id, "修正路1号")
        corrected = service.reparse(first.address_id, confirm_corrected=True)
        assert geo.calls == 2
        assert corrected.query_address == "修正路1号"
        assert corrected.confirmation_status == AddressConfirmationStatus.CORRECTED_AND_CONFIRMED.value
        assert corrected.geocode_version != first.geocode_version


def test_processor_reuses_confirmed_address_and_route_across_rows(tmp_path):
    geo, router = CountingGeocoder(), CountingRouter()
    rows = [row(2), row(99)]
    with CacheRepository(tmp_path / "cache.db") as repository:
        repository.create_task("task", "SHA", {}, "driving_real")
        outcomes = BusinessProcessor(repository, geocoder=geo, router=router).process(
            rows, task_id="task", mode=TaskMode.DRIVING_REAL,
            city_catalog=["广州", "深圳"], force_all=True,
        )
        assert geo.calls == 2
        assert router.calls == 1
        assert outcomes[1].cache_reused is True
        assert "可信地址库复用" in outcomes[1].explanation


def test_city_conflict_manual_confirmation_is_reused_without_warning(tmp_path):
    geo, router = CountingGeocoder(), CountingRouter()
    conflict_row = row(origin_city="广州", origin="广东省深圳市测试区测试路1号")
    with CacheRepository(tmp_path / "cache.db") as repository:
        processor = BusinessProcessor(repository, geocoder=geo, router=router)
        first = processor.process_row(conflict_row, TaskMode.DRIVING_REAL, ["广州", "深圳"])
        assert first.status == "查询成功—地址冲突待确认"
        service = AddressBookService(repository, geocoder=geo)
        service.confirm(address_id_for("广州", conflict_row.origin_address))
        second = processor.process_row(conflict_row, TaskMode.DRIVING_REAL, ["广州", "深圳"])
        assert "冲突待确认" not in second.status
        assert not second.warning_fields


def test_invalid_address_blocks_route(tmp_path):
    geo, router = CountingGeocoder(), CountingRouter()
    with CacheRepository(tmp_path / "cache.db") as repository:
        service = AddressBookService(repository, geocoder=geo)
        record = service.resolve("广州", "测试路", city_conflict=False).record
        service.mark_invalid(record.address_id)
        outcome = BusinessProcessor(repository, geocoder=geo, router=router).process_row(
            row(), TaskMode.DRIVING_REAL, ["广州", "深圳"]
        )
        assert outcome.status == "地址无效"
        assert outcome.distance_km is None
        assert router.calls == 0


def test_batch_confirms_high_confidence_candidates(tmp_path):
    geo = CountingGeocoder()
    with CacheRepository(tmp_path / "cache.db") as repository:
        service = AddressBookService(repository, geocoder=geo)
        record = service.resolve("广州", "测试路", city_conflict=False).record
        repository.put_address(replace(
            record, confirmation_status=AddressConfirmationStatus.UNCONFIRMED.value
        ))
        assert service.batch_confirm_high_trust() == 1
        assert repository.get_address(record.address_id).confirmation_status == AddressConfirmationStatus.SYSTEM_TRUSTED.value


def test_coordinate_change_invalidates_route(tmp_path):
    geo = CountingGeocoder()
    with CacheRepository(tmp_path / "cache.db") as repository:
        service = AddressBookService(repository, geocoder=geo)
        record = service.resolve("广州", "测试路", city_conflict=False).record
        repository.put_route(
            RouteResult("route", 1.0, mode="driving_real"), "o", "d",
            destination_address_id=record.address_id,
            destination_geocode_version=record.geocode_version,
        )
        service.update_coordinates(record.address_id, 113.5, 23.5)
        assert repository.get_route("route") is None


def test_pending_excel_export_import_matches_by_address_id(tmp_path):
    geo = CountingGeocoder(level="乡镇")
    with CacheRepository(tmp_path / "cache.db") as repository:
        service = AddressBookService(repository, geocoder=geo)
        record = service.resolve("广州", "测试镇仓库", city_conflict=False).record
        target = export_pending_addresses(tmp_path / "pending.xlsx", [record])
        workbook = load_workbook(target)
        sheet = workbook["待确认地址"]
        assert sheet["A2"].value == record.address_id
        sheet["G2"] = "是"
        sheet["H2"] = "现场已核实"
        workbook.save(target)
        workbook.close()
        summary = import_confirmation_results(target, service)
        assert (summary.success, summary.skipped, summary.failed) == (1, 0, 0)
        confirmed = repository.get_address(record.address_id)
        assert confirmed.confirmation_status == AddressConfirmationStatus.MANUALLY_CONFIRMED.value
        assert confirmed.confirmation_note == "现场已核实"


def test_import_rejects_duplicate_conflicts(tmp_path):
    geo = CountingGeocoder(level="乡镇")
    with CacheRepository(tmp_path / "cache.db") as repository:
        service = AddressBookService(repository, geocoder=geo)
        record = service.resolve("广州", "测试镇仓库", city_conflict=False).record
        target = export_pending_addresses(tmp_path / "pending.xlsx", [record])
        workbook = load_workbook(target)
        sheet = workbook["待确认地址"]
        sheet.append([record.address_id, "广州", "测试镇仓库", "", "乡镇", "另一个地址", "是", "冲突"])
        workbook.save(target)
        workbook.close()
        summary = import_confirmation_results(target, service)
        assert summary.failed == 1
        assert "重复且内容冲突" in summary.errors[0]
